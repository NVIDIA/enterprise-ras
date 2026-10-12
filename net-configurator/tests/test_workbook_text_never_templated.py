# SPDX-FileCopyrightText: Copyright (c) 2026 NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: MIT
"""B1 / GitLab #88 (ADR-0090): workbook text is data, never an Ansible template.

A banner or Custom_Config cell holding ``{{ lookup('pipe', ...) }}`` used to run a
command on the controller during ``make generate``. Two independent defences, so
two independent tests — a validation rejection must not hide a broken writer:

* ``TestValidationWorkflow`` — the normal path refuses the workbook, names the
  sheet and cell, writes no inventory, and nothing executes.
* ``TestInventoryWriterBypass`` — skips validation on purpose, then renders the
  real configs through Ansible: the payload stays literal (``!unsafe``), nothing
  executes, and tool-authored expressions (``{{ ldap_admin_password }}``) still
  resolve.
"""
import shutil
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

NC = Path(__file__).resolve().parent.parent
ARCH = "2-4-3-200"
_COPY_IGNORE = shutil.ignore_patterns(
    "output", "archive", ".git", ".venv*", "__pycache__", "*.pyc",
    ".pytest_cache", ".era-secrets", "legacy",
)
needs_ansible = pytest.mark.skipif(not shutil.which("ansible-playbook"),
                                   reason="ansible-playbook required")


def _payload(marker):
    return "{{ lookup('pipe', 'touch " + str(marker) + "') }}"


def _tree(tmp_path):
    dst = tmp_path / "nc"
    shutil.copytree(NC, dst, ignore=_COPY_IGNORE)
    return dst, dst / "input" / ARCH / "default" / f"{ARCH}.xlsx"


def _set_banner(xlsx, text):
    """Put `text` in Deployment's pre_login_message value cell; return its coordinate."""
    wb = openpyxl.load_workbook(xlsx)
    ws = wb["Deployment"]
    for row in ws.iter_rows():
        if row[0].value == "pre_login_message":
            row[1].value = text
            wb.save(xlsx)
            return row[1].coordinate
    raise AssertionError("Deployment sheet has no pre_login_message row")


def _add_custom_config(xlsx, text):
    wb = openpyxl.load_workbook(xlsx)
    ws = wb.create_sheet("Custom_Config")
    ws.append(["Switch_Location", "Config"])
    ws.append(["Function: core", text])
    wb.save(xlsx)
    return "B2"


def _run(cmd, cwd):
    return subprocess.run(cmd, cwd=cwd, capture_output=True, text=True)


class TestValidationWorkflow:
    @pytest.mark.parametrize("sheet,inject", [
        ("Deployment", _set_banner),
        ("Custom_Config", _add_custom_config),
    ])
    def test_malicious_cell_is_refused_with_sheet_and_cell(self, tmp_path, sheet, inject):
        marker = tmp_path / "PWN"
        dst, xlsx = _tree(tmp_path)
        cell = inject(xlsx, _payload(marker))

        v = _run([sys.executable, "scripts/validate_excel.py", str(xlsx)], dst)
        assert v.returncode == 1, v.stdout[-2000:]
        assert f"[{sheet}] cell {cell}: contains the template marker '{{{{'" in v.stdout

        p = _run([sys.executable, "scripts/excel_parser.py", "--arch", ARCH,
                  "--site", "default"], dst)
        assert p.returncode != 0
        assert not (dst / "output" / ARCH / "default" / "inventory").exists()
        assert not marker.exists()

    @pytest.mark.parametrize("text", ["{% if true %}x{% endif %}", "banner {# c #}"])
    def test_every_marker_kind_is_refused(self, tmp_path, text):
        dst, xlsx = _tree(tmp_path)
        cell = _set_banner(xlsx, text)
        v = _run([sys.executable, "scripts/validate_excel.py", str(xlsx)], dst)
        assert v.returncode == 1
        assert f"[Deployment] cell {cell}: contains the template marker" in v.stdout

    def test_single_brace_banner_placeholders_still_validate(self, tmp_path):
        dst, xlsx = _tree(tmp_path)
        _set_banner(xlsx, "Welcome to {hostname} at {site} ({arch})")
        v = _run([sys.executable, "scripts/validate_excel.py", str(xlsx)], dst)
        assert v.returncode == 0, v.stdout[-2000:]


@needs_ansible
class TestInventoryWriterBypass:
    def test_payload_stays_literal_through_ansible_render(self, tmp_path):
        marker = tmp_path / "PWN"
        dst, xlsx = _tree(tmp_path)
        _set_banner(xlsx, _payload(marker))
        inv = dst / "output" / ARCH / "default" / "inventory"

        p = _run([sys.executable, "scripts/excel_parser.py", "--arch", ARCH,
                  "--site", "default", "--skip-validate"], dst)
        assert p.returncode == 0, p.stderr[-2000:]

        gen = _run(["ansible-playbook", "playbooks/generate-cli-configs.yml",
                    "-i", f"output/{ARCH}/default/inventory/hosts",
                    "-e", f"config_output_dir=../output/{ARCH}/default/configs"], dst)
        assert gen.returncode == 0, gen.stderr[-2000:]
        assert not marker.exists(), "workbook text was executed as a Jinja template"
        written = (inv / "group_vars" / "all" / "main.yml").read_text()
        assert "pre_login_message: !unsafe" in written

        cfg = (dst / "output" / ARCH / "default" / "configs" / "core-01-config.sh").read_text()
        banner = [ln for ln in cfg.splitlines() if "system message pre-login" in ln]
        assert banner and "{{ lookup(" in banner[0] and str(marker) in banner[0]

    def test_tool_authored_expressions_still_resolve(self, tmp_path):
        dst, _ = _tree(tmp_path)
        p = _run([sys.executable, "scripts/excel_parser.py", "--arch", ARCH,
                  "--site", "default", "--skip-validate"], dst)
        assert p.returncode == 0, p.stderr[-2000:]
        play = dst / "playbooks" / "_probe_tool_exprs.yml"
        play.write_text(
            "- hosts: switches\n  gather_facts: false\n  connection: local\n"
            "  tasks:\n    - assert:\n        that:\n"
            "          - ldap.admin_password == ldap_admin_password\n"
            "          - ldap.admin_password is not search('{{ \"{{\" }}')\n"
            "      run_once: true\n")
        r = _run(["ansible-playbook", str(play), "-i",
                  f"output/{ARCH}/default/inventory/hosts"], dst)
        assert r.returncode == 0, r.stdout[-2000:]

    def test_hand_edited_custom_config_is_refused_at_render(self, tmp_path):
        """H1 (GitLab #89): a custom_config line that never went through the
        parser (hand-edited host_vars) fails the Ansible render; the shell text
        never reaches a config and nothing executes."""
        marker = tmp_path / "PWN"
        dst, _ = _tree(tmp_path)
        p = _run([sys.executable, "scripts/excel_parser.py", "--arch", ARCH,
                  "--site", "default", "--skip-validate"], dst)
        assert p.returncode == 0, p.stderr[-2000:]
        hv = dst / "output" / ARCH / "default" / "inventory" / "host_vars" / "core-01.yml"
        hv.write_text(hv.read_text() + f"custom_config:\n- 'nv set x; touch {marker}'\n")
        gen = _run(["ansible-playbook", "playbooks/generate-cli-configs.yml",
                    "-i", f"output/{ARCH}/default/inventory/hosts",
                    "-e", f"config_output_dir=../output/{ARCH}/default/configs"], dst)
        assert gen.returncode != 0
        assert "outside_the_nv_set_unset_grammar" in gen.stdout + gen.stderr
        assert not marker.exists()
        cfg = dst / "output" / ARCH / "default" / "configs" / "core-01-config.sh"
        assert not cfg.exists() or "touch" not in cfg.read_text()
