# SPDX-FileCopyrightText: Copyright (c) 2026 NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: MIT
"""
GitLab #85 / ERA-116 — an unrecognised Switch Role must fail, not become a server.

`utils.is_switch()` / `classify_node()` return False / `unknown` for a role
nobody recognises. Pre-fix the topology generator then took the server path
for that node everywhere — `_resolve_os` picked the Ubuntu image, port
allocation handed out ethN, links and layout treated it as a host — and
generation succeeded. The hostname fallback hides this on shipped workbooks
(`core-01` classifies by name); it bites on OEM-named switches, where the
Function cell is the only role source.

ADR-0057: role is declared by function, and a classification miss fails
loudly. Both the validator and the generator must refuse.

The mutation renames one core switch to an OEM-style name AND sets its
Function to `cor`; either change alone is rescued by the hostname fallback.
Pre-fix (recorded against the pre-change checkout): the validator emitted no
role error and the generator built the topology with the switch as a host.
"""
import os
import shutil
import sys
import uuid
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).parent))
sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
import workbook_access  # noqa: E402
from topology_generator import TopologyGenerator, UnclassifiedPeerRoleError  # noqa: E402
from validate_excel import validate_excel  # noqa: E402

REPO = Path(__file__).resolve().parent.parent
ARCH = "2-8-5-200"
BASE = REPO / "input" / ARCH / "default" / f"{ARCH}.xlsx"
OEM_NAME = "leaf-a1"
BOGUS_ROLE = "cor"

pytestmark = pytest.mark.skipif(
    not BASE.exists(), reason=f"default {ARCH} Excel not present")


def _misspell_one_switch(wb, new_role=BOGUS_ROLE, new_name=OEM_NAME):
    ws, cols, first = workbook_access.sheet_table(wb, "Switch Data")
    for r in range(first, ws.max_row + 1):
        if str(ws.cell(r, cols["function"]).value or "").strip() == "core":
            old = str(ws.cell(r, cols["name"]).value).strip()
            ws.cell(r, cols["function"]).value = new_role
            break
    else:
        raise AssertionError("no core switch found to mutate")
    assert workbook_access.rename_node(wb, old, new_name) > 1
    return old


def _make_site(prefix, mutate):
    """Scratch site under the repo (the generator writes output/<arch>/<site>).

    The site name is unique per process and per call so xdist workers never
    share a workbook or an output dir; both are pre-cleaned and removed after.
    """
    name = f"{prefix}-{os.getpid()}-{uuid.uuid4().hex[:8]}"
    in_dir = REPO / "input" / ARCH / name
    out_dir = REPO / "output" / ARCH / name
    shutil.rmtree(in_dir, ignore_errors=True)
    shutil.rmtree(out_dir, ignore_errors=True)
    in_dir.mkdir(parents=True)
    dst = in_dir / f"{ARCH}.xlsx"
    shutil.copy2(BASE, dst)
    wb = openpyxl.load_workbook(dst)
    mutate(wb)
    wb.save(dst)
    return name, dst, in_dir, out_dir


@pytest.fixture
def site():
    """Mutated workbook with an OEM-named switch whose Function is `cor`."""
    name, dst, in_dir, out_dir = _make_site("gl85-unknown-role", _misspell_one_switch)
    yield name, dst
    shutil.rmtree(in_dir, ignore_errors=True)
    shutil.rmtree(out_dir, ignore_errors=True)


@pytest.fixture
def ext_site():
    """Same misspelling, but the peer is renamed to an `ext-*` sentinel."""
    name, dst, in_dir, out_dir = _make_site(
        "gl85-ext-sentinel",
        lambda wb: _misspell_one_switch(wb, new_name="ext-fw-02"))
    yield name, dst
    shutil.rmtree(in_dir, ignore_errors=True)
    shutil.rmtree(out_dir, ignore_errors=True)


def test_validator_errors_on_unclassified_b_side(site):
    _, xlsx = site
    result = validate_excel(str(xlsx))
    hits = [e for e in result.errors
            if OEM_NAME in e and BOGUS_ROLE in e and "not a recognised role" in e]
    assert len(hits) == 1, (
        f"expected exactly one role error for {OEM_NAME}; errors={result.errors}")
    assert "treat this node as a server" in hits[0]


def test_generator_refuses_rather_than_rendering_a_server(site):
    name, _ = site
    with pytest.raises(UnclassifiedPeerRoleError) as exc:
        TopologyGenerator(REPO / "input" / ARCH / name / f"{ARCH}.xlsx", ARCH, name)
    msg = str(exc.value)
    assert OEM_NAME in msg and f"'{BOGUS_ROLE}'" in msg
    assert "refusing to render these as servers" in msg


def test_generator_does_not_write_a_topology(site):
    name, _ = site
    out = REPO / "output" / ARCH / name / "topology"
    with pytest.raises(UnclassifiedPeerRoleError):
        TopologyGenerator(REPO / "input" / ARCH / name / f"{ARCH}.xlsx", ARCH, name)
    assert not out.exists(), "generator wrote artifacts despite the role error"


def test_ext_sentinel_peer_is_skipped_by_both_validator_and_generator(ext_site):
    """ext-* B-sides are unmanaged gear (ADR-0047): neither side may fail on them."""
    name, xlsx = ext_site
    result = validate_excel(str(xlsx))
    assert not [e for e in result.errors if "not a recognised role" in e], result.errors
    TopologyGenerator(REPO / "input" / ARCH / name / f"{ARCH}.xlsx", ARCH, name)


def test_validator_resolves_nodes_function_before_wire_map_role():
    """Order must match parse_wiremap_excel._resolve: canonical Nodes Function
    wins over the Switch Role cell, so a stale 'cor' cell on a switch the Nodes
    tab declares `core` is not a role error (the generator builds it as core)."""
    from validate_excel import validate_wire_map, ValidationResult
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Wire Map"
    ws.append(["System Role", "System Name", "NIC Port", "Network Profile",
               "Switch Role", "Switch Name", "Switch Port"])
    ws.append(["gpu", "node-01", "eth0", "CPU/In-Band Network",
               "cor", OEM_NAME, "swp1"])
    result = ValidationResult()
    validate_wire_map(ws, result, nodes_function_map={OEM_NAME: "core"})
    assert not [e for e in result.errors if "not a recognised role" in e], result.errors
    result = ValidationResult()
    validate_wire_map(ws, result, nodes_function_map={})
    assert [e for e in result.errors if OEM_NAME in e and "'cor'" in e], result.errors


@pytest.mark.parametrize("role", ["NA", "na", "OUTBOUND", "outbound"])
def test_validator_skip_set_matches_generator(role):
    """The generator skips B-sides whose role is NA or OUTBOUND (dhcp-oob:eth0 ->
    outbound). The validator is the earlier, friendlier gate and must not reject
    a row the generator would happily build: `outbound` rows leave validate_wire_map
    at the 'outbound link' skip before the role guard runs; `NA` is excluded in the
    guard itself."""
    from validate_excel import validate_wire_map, ValidationResult
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Wire Map"
    ws.append(["System Role", "System Name", "NIC Port", "Network Profile",
               "Switch Role", "Switch Name", "Switch Port"])
    ws.append(["dhcp-oob", "dhcp-oob", "eth0", "OOB Management Network",
               role, "outbound", "eth0"])
    result = ValidationResult()
    validate_wire_map(ws, result, nodes_function_map={})
    assert not [e for e in result.errors if "not a recognised role" in e], result.errors


@pytest.mark.parametrize("role", ["core", "cs", "cl", "csl", "gs-plane1",
                                  "gl-plane2", "gsl-plane1", "oob-switch"])
def test_every_canonical_switch_role_still_classifies(role):
    from utils import classify_node, is_switch
    assert is_switch(role), role
    assert classify_node(role) != "unknown", role


def test_shipped_default_workbook_has_no_role_errors():
    result = validate_excel(str(BASE))
    assert not [e for e in result.errors if "not a recognised role" in e]
    TopologyGenerator(BASE, ARCH, "default")
