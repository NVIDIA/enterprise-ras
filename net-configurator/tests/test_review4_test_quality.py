# SPDX-FileCopyrightText: Copyright (c) 2026 NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: MIT
"""Review #4 F11 / F15 / F22 (GitLab #106): the suite itself stays honest.

- F11: `make lint` lints `tests/` too, so `.flake8`'s enforced F401 covers it.
- F15: an air-deploy.py import failure fails its tests instead of skipping them;
  `test_pre_public_polish`'s parameter sets are not empty; no test is skipped
  unconditionally (the retired legacy-spine comparisons were).
- F22: the shipped flat workbooks are in the zero-warning gate (their dual-plane
  loopbacks no longer warn). Private-tree generation is checked in
  `test_complete_workbook.py` itself.
"""
import ast
import re
import shutil
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest
import yaml

NC = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(NC / "scripts"))

from validate_excel import ValidationResult, validate_loopbacks  # noqa: E402


def _recipe(target):
    text = (NC / "Makefile").read_text()
    m = re.search(rf"^{target}:.*\n((?:\t.*\n|\n)+)", text, re.M)
    assert m, f"no {target} target"
    return m.group(1)


def test_make_lint_covers_tests():
    flake8 = [ln for ln in _recipe("lint").splitlines() if re.match(r"\s*flake8 ", ln)]
    assert flake8, "lint recipe runs no flake8"
    assert all(re.search(r"\bflake8 scripts/ tests/", ln) for ln in flake8), flake8


ARCHS = yaml.safe_load((NC / "scripts" / "valid_archs.yml").read_text())["archs"]
_AIR_DEPLOY_CLASSES = ("TestResolveInfraNodes or TestSwitchNIPasswordNotLeaked or "
                       "TestSwitchNIApplyVerifiesAppliedState or TestSwitchNIApplySilentFailureIsCaught")


def test_air_deploy_import_failure_is_not_a_skip(tmp_path):
    """Break air-deploy.py in a private copy: its tests must fail, not skip."""
    tree = tmp_path / "nc"
    shutil.copytree(NC, tree, ignore=shutil.ignore_patterns(
        "output", "archive", ".git", ".venv*", "__pycache__", "*.pyc", ".pytest_cache"))
    (tree / "scripts" / "air-deploy.py").write_text('raise ImportError("broken on purpose")\n')
    r = subprocess.run([sys.executable, "-m", "pytest", "tests/test_topology_generator.py",
                        "-q", "-p", "no:cacheprovider", "-o", "addopts=", "-rs",
                        "-k", _AIR_DEPLOY_CLASSES],
                       cwd=tree, capture_output=True, text=True)
    tail = r.stdout[-3000:]
    assert r.returncode == 1, tail
    m = re.search(r"(\d+) failed", tail)
    assert m and int(m.group(1)) >= 4, tail
    assert "skipped" not in tail.splitlines()[-1], tail


def _collected(test_file):
    r = subprocess.run([sys.executable, "-m", "pytest", "--collect-only", "-q",
                        "-p", "no:cacheprovider", "-o", "addopts=", test_file],
                       cwd=NC, capture_output=True, text=True)
    assert r.returncode == 0, r.stdout[-2000:] + r.stderr[-2000:]
    return r.stdout


@pytest.mark.parametrize("func", [
    "test_source_inventory_uses_underscore_oob_group",
    "test_inventory_children_reference_defined_groups",
    "test_inventory_no_hyphenated_group_headers",
])
def test_pre_public_polish_parameter_sets_are_not_empty(func):
    out = _collected("tests/test_pre_public_polish.py")
    ids = [ln for ln in out.splitlines() if f"::{func}[" in ln]
    assert "NOTSET" not in out and len(ids) >= 2 * len(ARCHS), (func, len(ids))


def _is_unconditional_skip(node):
    target = node.func if isinstance(node, ast.Call) else node
    return (isinstance(target, ast.Attribute) and target.attr == "skip"
            and isinstance(target.value, ast.Attribute) and target.value.attr == "mark")


def test_no_test_is_skipped_unconditionally():
    found = []
    for path in sorted((NC / "tests").glob("*.py")):
        for node in ast.walk(ast.parse(path.read_text(), str(path))):
            if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)):
                found += [f"{path.name}:{node.lineno} {node.name}"
                          for d in node.decorator_list if _is_unconditional_skip(d)]
            elif isinstance(node, ast.Assign) and any(
                    isinstance(tg, ast.Name) and tg.id == "pytestmark" for tg in node.targets):
                vals = node.value.elts if isinstance(node.value, (ast.List, ast.Tuple)) else [node.value]
                found += [f"{path.name}:{node.lineno} pytestmark" for v in vals if _is_unconditional_skip(v)]
    assert not found, found


# --- loopback_base warning -------------------------------------------------

def _loopbacks(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Loopbacks"
    for c, h in enumerate(["Switch", "Default", "OOB", "INBAND", "EXIT", "GPU"], start=1):
        ws.cell(row=1, column=c, value=h)
    for r, row in enumerate(rows, start=2):
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=v)
    return ws


def _base_warnings(rows):
    r = ValidationResult()
    nodes = [{"name": row[0], "function": "x"} for row in rows]
    validate_loopbacks(_loopbacks(rows), nodes, parsed_vlans=[],
                       settings={"loopback_base": "172.16.1"}, result=r)
    return [w for w in r.warnings if "outside Settings.loopback_base" in str(w)]


def test_plane_switches_in_their_own_plane_block_do_not_warn():
    assert _base_warnings([
        ["gl-plane1-01", "10.1.1.1/32", "", "", "", ""],
        ["gs-plane2-01", "10.2.1.41/32", "", "", "", ""],
        ["gsl-plane1-01", "10.9.9.9/32", "", "", "", ""],
        ["core-01", "172.16.1.1/32", "", "", "", ""],
    ]) == []


@pytest.mark.parametrize("name,ip", [
    ("gl-plane1-01", "10.2.1.1/32"),   # another plane's block
    ("leaf-01", "10.1.1.1/32"),        # not a plane switch
    ("gl-plane1-01", "10.1.2.1/32"),   # plane prefix, wrong /24
])
def test_anything_else_outside_loopback_base_still_warns(name, ip):
    warnings = _base_warnings([[name, ip, "", "", "", ""]])
    assert len(warnings) == 1 and name in str(warnings[0]), warnings


def test_shipped_gate_covers_every_flat_workbook():
    sys.path.insert(0, str(NC / "tests"))
    from test_shipped_workbooks_are_clean import FLAT
    assert {p.name for p in FLAT} == {f"{pre}-{a}.xlsx" for pre in ("sample", "largescale")
                                      for a in ARCHS}
