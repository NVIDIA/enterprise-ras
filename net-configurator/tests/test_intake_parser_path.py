# SPDX-FileCopyrightText: Copyright (c) 2026 NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: MIT
"""An intake-form workbook must parse to the same settings a legacy one does."""
import subprocess
import sys
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]

sys.path.insert(0, str(Path(__file__).parent))
from conftest import LEGACY_SAMPLE  # noqa: E402

# The frozen legacy sample -- see conftest. Half these tests are about telling
# the two formats apart, so one of each is needed and `input/` now holds only
# the intake form.
LEGACY_SRC = str(LEGACY_SAMPLE)
sys.path.insert(0, str(REPO / "scripts"))

import openpyxl


from excel_parser import is_intake_form, parse_nodes, process_excel_template


def test_intake_form_is_detected_by_sheet_names(tmp_path):
    src = LEGACY_SRC
    dst = tmp_path / "intake.xlsx"
    subprocess.run(["python3", "scripts/convert_workbook.py", src, str(dst)],
                   cwd=REPO, check=True)
    assert is_intake_form(openpyxl.load_workbook(dst)) is True


def test_legacy_workbook_is_not_detected_as_intake_form():
    src = LEGACY_SRC
    assert is_intake_form(openpyxl.load_workbook(src)) is False


def test_legacy_workbook_with_a_stray_wiremap_and_deployment_tab_still_takes_the_legacy_path(tmp_path):
    # is_intake_form() requires both intake markers present AND every
    # legacy-required sheet absent, so a legacy workbook that picks up a
    # stray Wiremap/Deployment tab (scratch notes, a half-finished manual
    # conversion, ...) cannot flip onto the intake path and hit its R4-05
    # refusal on sheets it doesn't actually have in intake form.
    src = LEGACY_SRC
    wb = openpyxl.load_workbook(src)
    wb.create_sheet("Wiremap")
    wb.create_sheet("Deployment")
    mixed = tmp_path / "legacy-with-stray-intake-tabs.xlsx"
    wb.save(mixed)

    assert is_intake_form(openpyxl.load_workbook(mixed)) is False


def test_converted_workbook_yields_the_same_node_count_as_the_legacy_original(tmp_path):
    # Cheapest early signal that Task 10's byte-identical assertion has a
    # chance: if the intake path drops or duplicates nodes relative to the
    # legacy Nodes sheet, nothing downstream can possibly match.
    src = LEGACY_SRC
    dst = tmp_path / "intake.xlsx"
    subprocess.run(["python3", "scripts/convert_workbook.py", src, str(dst)],
                   cwd=REPO, check=True)

    from excel_parser import _parse_intake_nodes

    legacy_wb = openpyxl.load_workbook(src, data_only=True)
    legacy_nodes = parse_nodes(legacy_wb['Nodes'])

    intake_wb = openpyxl.load_workbook(dst, data_only=True)
    intake_nodes = _parse_intake_nodes(intake_wb)

    assert len(intake_nodes) == len(legacy_nodes)


def test_process_excel_template_generates_from_an_intake_workbook(tmp_path):
    """An intake workbook must GENERATE, not be refused.

    Phase 0 refused it outright: everything below the settings/nodes branch
    point read legacy-only sheets, so letting it through produced a raw
    `KeyError`. Phase 1 builds those sheets in memory (`intake_to_legacy`),
    so the pipeline runs. This asserts the refusal is gone AND that real
    output lands -- a test that only checked "no exception" would pass
    against a function that silently did nothing.
    """
    dst = tmp_path / "intake.xlsx"
    subprocess.run(["python3", "scripts/convert_workbook.py", LEGACY_SRC,
                    str(dst)], cwd=REPO, check=True)

    out = tmp_path / "out"
    out.mkdir()
    process_excel_template(str(dst), out)

    hosts = out / "hosts"
    assert hosts.is_file(), f"no inventory written to {out}"
    assert hosts.read_text().strip(), "inventory file is empty"
    assert list((out / "host_vars").glob("*.yml")), "no host_vars generated"


def test_validate_excel_accepts_an_intake_workbook(tmp_path):
    """`validate_excel` must validate an intake workbook, not reject it.

    `make generate` gates on this validator at step 0, so a rejection here
    makes the whole format unusable regardless of what the parser can do.
    Phase 0 rejected it (correctly, since generation did not work); Phase 1
    builds the legacy-shaped views first, so every existing validator runs
    against the same content it always did.

    Asserts specifically that no error mentions a missing legacy sheet --
    that was the symptom when the validator did not understand the format.
    """
    import sys
    from pathlib import Path
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
    import validate_excel
    from convert_workbook import convert

    intake = tmp_path / "intake.xlsx"
    convert(LEGACY_SRC, str(intake))

    result = validate_excel.validate_excel(str(intake))
    blob = " ".join(str(e) for e in result.errors)

    assert "Missing required sheet" not in blob, (
        f"validator treated an intake workbook as a broken legacy one: {blob}")
    assert not result.errors, f"intake workbook failed validation: {blob}"
