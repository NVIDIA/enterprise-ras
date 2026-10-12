#!/usr/bin/env python3
# SPDX-FileCopyrightText: Copyright (c) 2026 NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: MIT
"""`<arch>-complete.xlsx` — the authored workbook with every address filled in.

An intake form can be authored sparse and let the tool allocate. The completed
workbook is what makes that reviewable: it sits beside the source in `input/`
and shows exactly what was chosen.

The load-bearing property is that it is still a valid INPUT — feed it back and
you get the same configs, with nothing left to allocate.
"""
import shutil
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

NC = Path(__file__).resolve().parent.parent
ARCH = "2-8-9-800"
SRC = NC / "input" / ARCH / "default" / f"{ARCH}.xlsx"

sys.path.insert(0, str(NC / "scripts"))

# Generation writes into input/ and output/ (and `.era-context` at the tool
# root), so every test works in its own copy of the tool tree. Writing into the
# shared tree raced under `make test-fast` with any test that globs input/.
_TREE_IGNORE = shutil.ignore_patterns(
    "output", "archive", ".git", ".venv*", "__pycache__", "*.pyc",
    ".pytest_cache", ".era-secrets", ".era-context", "legacy", "tests")


def _tree(tmp_path):
    nc = tmp_path / "nc"
    shutil.copytree(NC, nc, ignore=_TREE_IGNORE)
    return nc


def _run(tree, *args):
    return subprocess.run([sys.executable, *args], cwd=tree,
                          capture_output=True, text=True)


def _site(tree, name):
    d = tree / "input" / ARCH / name
    d.mkdir(parents=True)
    return d


@pytest.fixture
def completed(tmp_path):
    """An intake site with its completed workbook written, in a private tree."""
    tree = _tree(tmp_path)
    site = "pytest-complete"
    d = _site(tree, site)
    shutil.copy2(SRC, d / f"{ARCH}.xlsx")

    emit = _run(tree, "scripts/emit_complete_workbook.py",
                "--source", str(d / f"{ARCH}.xlsx"),
                "--arch", ARCH, "--site", site)
    assert emit.returncode == 0, emit.stderr
    return d / f"{ARCH}-complete.xlsx", site


def test_generation_stays_in_the_private_tree(completed, tmp_path):
    path, site = completed
    assert path.is_relative_to(tmp_path), path
    for shared in (NC / "input" / ARCH / site, NC / "output" / ARCH / site):
        assert not shared.exists(), f"generation wrote the shared tree: {shared}"


def test_completed_workbook_sits_beside_the_source(completed):
    path, _site_name = completed
    assert path.is_file()
    assert (path.parent / f"{ARCH}.xlsx").is_file(), "source must be kept"


def test_switch_addresses_are_filled_in(completed):
    """Every switch gets its resolved addresses, not just some."""
    path, _ = completed
    ws = openpyxl.load_workbook(path)["Switch Data"]
    hdr = {str(ws.cell(2, c).value).strip().lower(): c
           for c in range(1, ws.max_column + 1) if ws.cell(2, c).value}

    for col in ("mgmt ip", "loopback ip", "asn"):
        assert col in hdr, f"{col!r} column missing: {sorted(hdr)}"

    rows = [r for r in range(3, ws.max_row + 1) if ws.cell(r, hdr["name"]).value]
    assert rows, "no switch rows -- test would be vacuous"
    for r in rows:
        name = ws.cell(r, hdr["name"]).value
        assert ws.cell(r, hdr["loopback ip"]).value, f"{name}: no loopback"
        assert ws.cell(r, hdr["mgmt ip"]).value, f"{name}: no mgmt IP"


def test_mgmt_ip_is_the_oob_address_not_the_air_plane(completed):
    """Regression: `ansible_host` is the Air-plane address, not `Mgmt IP`.

    Writing it produced a workbook that generated DIFFERENT configs when fed
    back in — caught only by the round-trip check below, never by inspection.
    """
    path, site = completed
    ws = openpyxl.load_workbook(path)["Switch Data"]
    hdr = {str(ws.cell(2, c).value).strip().lower(): c
           for c in range(1, ws.max_column + 1) if ws.cell(2, c).value}
    values = [str(ws.cell(r, hdr["mgmt ip"]).value or "")
              for r in range(3, ws.max_row + 1)
              if ws.cell(r, hdr["name"]).value]
    assert values, "no rows -- vacuous"
    assert not any(v.startswith("172.20.") for v in values), (
        f"Mgmt IP holds air-mgmt addresses instead of OOB ones: {values[:4]}")


def test_every_resolved_host_address_appears(completed):
    """Host Data expands to one row per interface, so bond and GPU IPs show."""
    path, _ = completed
    ws = openpyxl.load_workbook(path)["Host Data"]
    hdr = {str(ws.cell(1, c).value).strip().lower(): c
           for c in range(1, ws.max_column + 1) if ws.cell(1, c).value}
    ips = {str(ws.cell(r, hdr["ip"]).value or "")
           for r in range(2, ws.max_row + 1)}

    assert any(v.startswith("192.168.200.") for v in ips), "no OOB addresses"
    assert any(v.startswith("172.16.") for v in ips), "no bond addresses"
    assert any(v.startswith("192.168.1.") or v.startswith("192.168.2.")
               for v in ips), "no GPU-plane addresses"


def test_completed_workbook_round_trips(completed):
    """THE property: re-import it and nothing changes.

    This is what makes the completed workbook an input rather than a report.
    It caught two real bugs -- the Air-plane mgmt address, and Host Data
    expanding to one row per interface in a way the node reader could not
    read back.
    """
    path, site = completed
    tree = path.parents[3]  # <tree>/input/<arch>/<site>/<file>
    rt = "pytest-roundtrip"
    d = _site(tree, rt)
    shutil.copy(path, d / f"{ARCH}.xlsx")
    gen = _run(tree, "scripts/excel_parser.py", "--arch", ARCH, "--site", rt,
               "--skip-validate")
    assert gen.returncode == 0, f"completed workbook is not a valid input:\n{gen.stderr}"

    original = tree / "output" / ARCH / site / "inventory"
    regenerated = tree / "output" / ARCH / rt / "inventory"
    files = {p.relative_to(original) for p in original.rglob("*") if p.is_file()}
    assert files, "no inventory generated -- vacuous"

    # `common.site` is the resolved SITE (ADR-0086); the two legs generate
    # into different sites, so exactly that one line in all/main.yml differs.
    def _norm(p, rel, site_name):
        data = p.read_bytes()
        if str(rel) != "group_vars/all/main.yml":
            return data
        needle = f"  site: {site_name}\n".encode()
        assert data.count(needle) == 1, f"{rel}: expected one common.site line"
        return data.replace(needle, b"  site: <SITE>\n")

    differing = [str(rel) for rel in files
                 if _norm(original / rel, rel, site)
                 != _norm(regenerated / rel, rel, rt)]
    assert not differing, (
        "re-importing the completed workbook changed the output in "
        f"{len(differing)} file(s): {differing[:8]}")


def test_host_data_keeps_its_position_in_the_tab_order(completed):
    """Rewriting the sheet must not shove it to the end of the tab bar."""
    path, _ = completed
    names = openpyxl.load_workbook(path).sheetnames
    assert names.index("Host Data") < names.index("Deployment"), (
        f"Host Data moved out of template order: {names}")


def test_host_rows_use_the_templates_column_meanings(completed):
    """Physical Interface = the Wire Map NIC label; NICOSName = the OS name.

    This is what makes a Host Data row traceable back to a cable. An earlier
    version put `ethN` in Physical Interface and left NICOSName blank, so the
    sheet could not be joined to the wiremap at all.
    """
    path, _ = completed
    ws = openpyxl.load_workbook(path)["Host Data"]
    hdr = {str(ws.cell(1, c).value).strip().lower(): c
           for c in range(1, ws.max_column + 1) if ws.cell(1, c).value}

    physical, osname = [], []
    for r in range(2, ws.max_row + 1):
        phys = ws.cell(r, hdr["physical interface"]).value
        nic = ws.cell(r, hdr["nicosname"]).value
        if phys:
            physical.append(str(phys))
        if nic:
            osname.append(str(nic))

    assert physical, "no Physical Interface values -- the wiremap join is dead"
    assert not any(v.startswith("eth") for v in physical), (
        f"Physical Interface holds OS names instead of NIC labels: {physical[:4]}")
    assert osname and all(v.startswith("eth") for v in osname), (
        f"NICOSName should hold ethN: {osname[:4]}")

    # Every Physical Interface must exist in the source Wire Map, or the join
    # is producing labels from somewhere it should not be.
    src = path.parent / path.name.replace("-complete.xlsx", ".xlsx")
    import sys as _sys
    _sys.path.insert(0, str(NC / "scripts"))
    import intake_to_legacy
    wm = intake_to_legacy.load_workbook(src, data_only=True)["Wire Map"]
    labels = {str(wm.cell(r, c).value) for r in range(2, wm.max_row + 1)
              for c in (3, 7)}
    unknown = sorted({v for v in physical if v not in labels})
    assert not unknown, f"NIC labels not present in the Wire Map: {unknown[:5]}"


def test_bond_members_are_identifiable(completed):
    """Bond composition must be readable: members share a Logical Interface."""
    path, _ = completed
    ws = openpyxl.load_workbook(path)["Host Data"]
    hdr = {str(ws.cell(1, c).value).strip().lower(): c
           for c in range(1, ws.max_column + 1) if ws.cell(1, c).value}

    bonds = {}
    for r in range(2, ws.max_row + 1):
        host = ws.cell(r, hdr["system name"]).value
        logical = str(ws.cell(r, hdr["logical interface"]).value or "")
        nic = ws.cell(r, hdr["nicosname"]).value
        if logical.startswith("bond") and host and nic:
            bonds.setdefault((str(host), logical), []).append(str(nic))

    assert bonds, "no bond rows found -- test would be vacuous"
    multi = [k for k, v in bonds.items() if len(v) > 1]
    assert multi, (
        "every bond shows a single member -- membership is not represented: "
        f"{dict(list(bonds.items())[:3])}")


def test_every_import_path_writes_the_completed_workbook(tmp_path):
    """The emit must live in import-excel.py, not in one Makefile target.

    Four targets shell out to this script -- `import`, `deploy`, and two
    others -- so a step hung on one of them is a step the other three
    silently skip. `make deploy` produced no completed workbook for exactly
    that reason. Asserts the script itself does it, which is what makes it
    true for every caller.
    """
    # import-excel.py also writes `.era-context` at the tool root; run it in a
    # private copy so other workers copying the tree never see that file churn.
    nc = tmp_path / "nc"
    shutil.copytree(NC, nc, ignore=shutil.ignore_patterns(
        "output", "archive", ".git", ".venv*", "__pycache__", "*.pyc",
        ".pytest_cache", ".era-secrets", ".era-context", "legacy", "tests"))
    site = "pytest-importpath"
    src_dir = nc / "input" / ARCH / site

    staged = tmp_path / f"{ARCH}.xlsx"
    shutil.copy2(SRC, staged)

    r = subprocess.run([sys.executable, "scripts/import-excel.py", str(staged),
                        "--site", site, "--yes"],
                       cwd=nc, capture_output=True, text=True)
    assert r.returncode == 0, f"import failed:\n{r.stdout}\n{r.stderr}"

    complete = src_dir / f"{ARCH}-complete.xlsx"
    assert complete.is_file(), (
        "import-excel.py did not write the completed workbook, so any "
        f"caller other than `make import` will not get one either.\n{r.stdout}")
    assert (src_dir / f"{ARCH}.xlsx").is_file(), "source workbook missing"
