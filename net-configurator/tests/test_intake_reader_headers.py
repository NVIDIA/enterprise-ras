# SPDX-FileCopyrightText: Copyright (c) 2026 NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: MIT
"""Locks in two-row grouped-header handling for the intake form."""
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).parent.parent / "scripts"))
sys.path.insert(0, str(Path(__file__).parent))

from intake_fixtures import build_wiremap_sheet
from intake_reader import column_map, header_row


def test_header_row_is_two_when_row_one_is_a_group_label():
    wb = openpyxl.Workbook()
    ws = build_wiremap_sheet(wb, [])
    assert header_row(ws) == 2


def test_column_map_normalises_case_and_whitespace():
    wb = openpyxl.Workbook()
    ws = build_wiremap_sheet(wb, [])
    cols = column_map(ws)
    assert cols["system name (a)"] == 4
    assert cols["system name (z)"] == 13
    assert cols["display in air"] == 23


def test_header_row_is_one_for_a_plain_single_row_header():
    # Host Data / Network Data rely on this: row 1 holds real column names
    # with no overlap at all against the known group labels.
    wb = openpyxl.Workbook()
    ws = wb.create_sheet("Host Data")
    for col, name in enumerate(["System Name", "Physical Interface", "IP"], start=1):
        ws.cell(row=1, column=col, value=name)
    assert header_row(ws) == 1


def test_header_row_raises_when_row_one_partially_overlaps_group_labels():
    # A template edit that adds one more band label than `_GROUP_LABELS`
    # knows about must not be silently read as a plain header -- that
    # would map the group-label row as column names and find nothing.
    wb = openpyxl.Workbook()
    ws = wb.create_sheet("Switch Data")
    for col, name in enumerate(["Switch", "Switch IP's", "L3", "L2", "Extra Band"], start=1):
        ws.cell(row=1, column=col, value=name)
    with pytest.raises(ValueError) as exc_info:
        header_row(ws)
    message = str(exc_info.value)
    assert "switch" in message.lower()
    assert "extra band" in message.lower()
