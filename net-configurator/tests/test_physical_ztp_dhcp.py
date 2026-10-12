# SPDX-FileCopyrightText: Copyright (c) 2026 NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: MIT
"""H5 / GitLab #93: physical (`Deploy in Air = No`) ZTP is supported in nc-v8.

Before this, a physical site was generated as if it were an Air sim, and its
ZTP server could not provision a real switch:

* dnsmasq rendered every static-IP switch as ``dhcp-host=<mac>,<host>,ignore``,
  so the server refused the very DISCOVER that starts ZTP;
* the inventory carried the Air infra nodes (``oob-server-01``,
  ``external-conn``) and an ``eth1`` air-mgmt (172.20.0.0/24) ZTP interface
  that does not exist off Air;
* in L3 OOB mode the ZTP server had no interface on the OOB subnet the
  switches actually DHCP on.

Each case generates a throwaway copy of the tool tree from a shipped workbook,
then renders the real dnsmasq templates with Ansible against the generated
inventory.
"""
import ipaddress
import re
import shutil
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest
import yaml

sys.path.insert(0, str(Path(__file__).parent))
import workbook_access  # noqa: E402

NC = Path(__file__).resolve().parent.parent
ARCH = "2-4-3-200"
AIR_ONLY_INFRA = {"oob-server-01", "external-conn"}
_COPY_IGNORE = shutil.ignore_patterns(
    "output", "archive", ".git", ".venv*", "__pycache__", "*.pyc",
    ".pytest_cache", ".era-secrets", "legacy",
)
ANSIBLE_PLAYBOOK = shutil.which("ansible-playbook") or str(Path(sys.executable).parent / "ansible-playbook")


def _set_ztp_mac(wb, name, mac):
    ws, cols, first = workbook_access.sheet_table(wb, "Switch Data")
    for r in range(first, ws.max_row + 1):
        if str(ws.cell(r, cols["name"]).value or "").strip() == name:
            ws.cell(r, cols["ztp mac address"]).value = mac
            return
    raise AssertionError(f"no Switch Data row for {name}")


def _groups(hosts_text):
    groups, current = {}, None
    for line in hosts_text.splitlines():
        line = line.strip()
        if line.startswith("[") and line.endswith("]"):
            current = line[1:-1]
            groups[current] = []
        elif line and current and not line.startswith("#") and ":" not in current:
            groups[current].append(line.split()[0])
    return groups


def _site(tmp_path, deploy_in_air, oob_mode):
    dst = tmp_path / "nc"
    shutil.copytree(NC, dst, ignore=_COPY_IGNORE)
    xlsx = dst / "input" / ARCH / "default" / f"{ARCH}.xlsx"
    wb = openpyxl.load_workbook(xlsx)
    workbook_access.set_setting(wb, "deploy_in_air", "Yes" if deploy_in_air else "No")
    workbook_access.set_setting(wb, "oob_uplink_mode", oob_mode)
    switches = {}
    for i, row in enumerate(r for r in workbook_access.node_rows(wb)
                            if r["is_switch"] and r.get("mgmt_ip")):
        mac = f"48:b0:2d:aa:01:{i + 1:02x}"
        _set_ztp_mac(wb, str(row["name"]), mac)
        switches[str(row["name"])] = (mac, str(row["mgmt_ip"]))
    wb.save(xlsx)
    r = subprocess.run([sys.executable, "scripts/excel_parser.py", "--arch", ARCH,
                        "--site", "default", "--skip-validate"],
                       cwd=dst, capture_output=True, text=True)
    assert r.returncode == 0, r.stderr[-2000:]
    inv = dst / "output" / ARCH / "default" / "inventory"
    return dst, inv, switches


def _render(tmp_path, dst, inv):
    out = tmp_path / "rendered"
    out.mkdir()
    tpl = dst / "roles" / "ztp-server" / "templates"
    play = tmp_path / "render.yml"
    play.write_text(yaml.safe_dump([{
        "hosts": "localhost", "gather_facts": False, "connection": "local",
        "tasks": [{"template": {"src": str(tpl / name), "dest": str(out / name[:-3])}}
                  for name in ("dnsmasq-ztp.conf.j2", "dnsmasq.conf.j2")],
    }]))
    r = subprocess.run([ANSIBLE_PLAYBOOK, "-i", str(inv / "hosts"), str(play)],
                       cwd=dst, capture_output=True, text=True, timeout=300)
    assert r.returncode == 0, r.stdout[-3000:] + r.stderr[-2000:]
    return (out / "dnsmasq-ztp.conf").read_text(), (out / "dnsmasq.conf").read_text()


@pytest.mark.parametrize("oob_mode", ["L3", "L2"])
def test_physical_site_reserves_each_switch_on_a_listening_ztp_subnet(tmp_path, oob_mode):
    dst, inv, switches = _site(tmp_path, deploy_in_air=False, oob_mode=oob_mode)
    main = yaml.safe_load((inv / "group_vars" / "all" / "main.yml").read_text())
    ifaces = main["ztp_interfaces"]
    assert ifaces, "a physical site needs at least one ZTP interface"
    air = ipaddress.ip_network(main["air_mgmt_subnet"])
    for i in ifaces:
        net = ipaddress.ip_network(i["network"])
        assert i["purpose"] != "air-mgmt" and not net.overlaps(air), i
        assert ipaddress.ip_address(i["ip"]) in net, i
        pool = (ipaddress.ip_address(i["dhcp_range_start"]), ipaddress.ip_address(i["dhcp_range_end"]))
        assert not pool[0] <= ipaddress.ip_address(i["ip"]) <= pool[1], \
            f"the ZTP server's own address {i['ip']} is inside its DHCP pool {pool}"
    assert main["air_mgmt_subnet"] not in main["ztp_allow_subnets"]
    assert main["ztp_server_ip"] == ifaces[0]["ip"], \
        "setup-ztp-server and the IP report must name the address DHCP option 239 serves"

    ztp_conf, base_conf = _render(tmp_path, dst, inv)
    assert ",ignore" not in ztp_conf
    listening = {m.group(1): m.group(2) for m in re.finditer(
        r"^dhcp-range=interface:(\S+?),(\S+?),", base_conf, re.M)}
    assert set(listening) == {i["name"] for i in ifaces}
    for i in ifaces:
        assert f"dhcp-option=interface:{i['name']},239,http://{i['ip']}/scripts/ztp.sh" in base_conf
    nets = [ipaddress.ip_network(i["network"]) for i in ifaces]
    for name, (mac, ip) in switches.items():
        hostname = yaml.safe_load((inv / "host_vars" / f"{name}.yml").read_text()).get("hostname", name)
        assert f"dhcp-host={mac},{hostname},{ip},12h" in ztp_conf, name
        assert any(ipaddress.ip_address(ip) in n for n in nets), \
            f"{name} {ip} is outside every ZTP subnet; its reservation could never be offered"


@pytest.mark.parametrize("oob_mode,kept", [("L3", {"external-dhcp", "utility"}), ("L2", {"dhcp-oob"})])
def test_physical_site_has_no_air_only_infra_nodes(tmp_path, oob_mode, kept):
    _, inv, _ = _site(tmp_path, deploy_in_air=False, oob_mode=oob_mode)
    groups = _groups((inv / "hosts").read_text())
    infra = set(groups.get("dhcp", [])) | set(groups.get("oob_server", [])) | set(groups.get("jump", []))
    assert infra == kept
    assert not groups.get("oob_server")
    for name in AIR_ONLY_INFRA:
        assert not (inv / "host_vars" / f"{name}.yml").exists(), name


def test_switching_a_site_to_physical_sweeps_its_air_only_host_vars(tmp_path):
    dst, inv, _ = _site(tmp_path, deploy_in_air=True, oob_mode="L2")
    assert (inv / "host_vars" / "oob-server-01.yml").exists()
    xlsx = dst / "input" / ARCH / "default" / f"{ARCH}.xlsx"
    wb = openpyxl.load_workbook(xlsx)
    workbook_access.set_setting(wb, "deploy_in_air", "No")
    wb.save(xlsx)
    r = subprocess.run([sys.executable, "scripts/excel_parser.py", "--arch", ARCH,
                        "--site", "default", "--skip-validate"],
                       cwd=dst, capture_output=True, text=True)
    assert r.returncode == 0, r.stderr[-2000:]
    assert not (inv / "host_vars" / "oob-server-01.yml").exists()


@pytest.mark.parametrize("oob_mode", ["L3", "L2"])
def test_air_site_keeps_its_air_mgmt_ztp_interface(tmp_path, oob_mode):
    _, inv, _ = _site(tmp_path, deploy_in_air=True, oob_mode=oob_mode)
    main = yaml.safe_load((inv / "group_vars" / "all" / "main.yml").read_text())
    assert main["ztp_interfaces"][0]["purpose"] == "air-mgmt"
    assert main["air_mgmt_subnet"] in main["ztp_allow_subnets"]
