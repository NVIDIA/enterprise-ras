# SPDX-FileCopyrightText: Copyright (c) 2026 NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: MIT
"""ERA-137: a switch that is Enabled=No must not leave dangling link endpoints.

Disabled Wire Map rows are kept so breakout analysis still sees the ports, and
``_build_nodes`` correctly drops the disabled switch. But the unconnected-stub
pass walked every switch named in the Wire Map, so ``gsl-plane2-02``
(Enabled=No) produced 97 ``[{node: gsl-plane2-02, ...}, "unconnected"]`` links
with no matching node, and ``make validate-topology`` passed anyway.
"""
import json
import sys
from pathlib import Path

import openpyxl
import pytest

SCRIPTS = Path(__file__).parent.parent / "scripts"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

from topology_generator import TopologyGenerator, TopologyValidator  # noqa: E402

WM_HEADER = [
    "Display in Air", "System Role", "System Name", "NIC/Port",
    "Speed", "Description", "Network Profile",
    "Disabled by Neighbor", "Speed2", "Description2",
    "Switch Role", "Switch Name", "Switch Port",
]


def _workbook(tmp_path, arch="2-4-3-200", core02_enabled="No"):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws_settings = wb.create_sheet("Settings")
    for row in (["VERSIONS", None], ["Switch Function", "Cumulus Version"],
                ["core", "5.18.0"], ["oob-switch", "5.18.0"]):
        ws_settings.append(row)
    ws_air = wb.create_sheet("Air_Only")
    ws_air.append(["Display in Air", "System Role", "System Name", "NIC/Port",
                   "Network Profile", "Switch Role", "Switch Name", "Switch Port"])
    ws_air.append([None] * 8)
    ws_air.append(["Friendly Version", "Air Image"])
    ws_air.append(["5.18.0", "cumulus-vx-5.18.0"])
    ws_nodes = wb.create_sheet("Nodes")
    ws_nodes.append(["Function", "Name", "MAC", "Mgmt IP", "Prefix", "Gateway", "ZTP", "Enabled"])
    ws_nodes.append(["gpu", "su-01-node-01", "", "192.168.200.10", "24", "192.168.200.1", "Yes", "Yes"])
    ws_nodes.append(["core", "core-01", "", "192.168.200.2", "24", "192.168.200.1", "Yes", "Yes"])
    ws_nodes.append(["core", "core-02", "", "192.168.200.3", "24", "192.168.200.1", "Yes", core02_enabled])
    ws_wm = wb.create_sheet("Wire Map")
    ws_wm.append(WM_HEADER)
    ws_wm.append(["Yes", "su-01-node-01", "su-01-node-01", "NIC1_P1", None, None,
                  "CPU/In-Band Network", None, None, None, "core-01", "core-01", "swp1s0"])
    ws_wm.append(["Yes", "su-01-node-01", "su-01-node-01", "NIC1_P2", None, None,
                  "CPU/In-Band Network", None, None, None, "core-02", "core-02", "swp1s0"])
    ws_wm.append(["Yes", "core", "core-01", "swp31", None, None,
                  "Peerlink", None, None, None, "core", "core-02", "swp31"])
    input_dir = tmp_path / "input" / arch / "default"
    input_dir.mkdir(parents=True)
    path = input_dir / f"{arch}.xlsx"
    wb.save(str(path))
    return path


def _orphans(topo):
    return TopologyGenerator.orphan_link_endpoints(topo["content"]["nodes"],
                                                   topo["content"]["links"])


def test_disabled_switch_gets_no_stubs_and_no_orphan_endpoints(tmp_path):
    topo = TopologyGenerator(_workbook(tmp_path), "2-4-3-200").generate()
    nodes = topo["content"]["nodes"]
    assert "core-01" in nodes and "core-02" not in nodes
    assert _orphans(topo) == []
    assert not any(isinstance(l[0], dict) and l[0]["node"] == "core-02"
                   for l in topo["content"]["links"])


def test_enabled_switch_still_gets_its_stubs(tmp_path):
    topo = TopologyGenerator(_workbook(tmp_path, core02_enabled="Yes"), "2-4-3-200").generate()
    assert "core-02" in topo["content"]["nodes"]
    assert _orphans(topo) == []
    core02_ifaces = {ep["interface"] for l in topo["content"]["links"]
                     for ep in l if isinstance(ep, dict) and ep["node"] == "core-02"}
    assert {"swp1s0", "swp31", "eth0"} <= core02_ifaces


def test_generator_refuses_a_topology_with_orphan_endpoints(tmp_path):
    gen = TopologyGenerator(_workbook(tmp_path), "2-4-3-200")
    with pytest.raises(ValueError, match="ghost-01"):
        gen._assert_link_endpoints_exist(
            {"core-01": {}}, [[{"node": "ghost-01", "interface": "swp1"}, "unconnected"]])


def test_validator_flags_orphan_link_endpoints(tmp_path, capsys):
    path = _workbook(tmp_path)
    topo = TopologyGenerator(path, "2-4-3-200").generate()
    topo["content"]["links"].append([{"node": "core-02", "interface": "swp5"}, "unconnected"])
    topo_path = tmp_path / "topology.json"
    topo_path.write_text(json.dumps(topo))
    validator = TopologyValidator(path, topo_path, "2-4-3-200")
    assert validator.validate() is False
    assert any("Orphan link endpoint: core-02" in e for e in validator.errors)


def test_validator_passes_the_disabled_switch_topology(tmp_path, capsys):
    path = _workbook(tmp_path)
    topo_path = tmp_path / "topology.json"
    topo_path.write_text(json.dumps(TopologyGenerator(path, "2-4-3-200").generate()))
    validator = TopologyValidator(path, topo_path, "2-4-3-200")
    assert validator.validate() is True, validator.errors
