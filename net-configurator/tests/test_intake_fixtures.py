# SPDX-FileCopyrightText: Copyright (c) 2026 NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: MIT
"""Tests for the intake-form fixture builders."""
import sys
from pathlib import Path

TESTS = Path(__file__).resolve().parent
if str(TESTS) not in sys.path:
    sys.path.insert(0, str(TESTS))

import openpyxl
from intake_fixtures import build_wiremap_sheet


def test_wiremap_fixture_has_grouped_two_row_header():
    wb = openpyxl.Workbook()
    ws = build_wiremap_sheet(wb, [])
    assert ws["A1"].value == "Source"
    assert ws["J1"].value == "Destination"
    assert ws["A2"].value == "DC Location (A)"
    assert ws["M2"].value == "System Name (Z)"
    assert ws["W2"].value == "Display in Air"
