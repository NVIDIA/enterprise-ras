# SPDX-FileCopyrightText: Copyright (c) 2026 NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: MIT
"""validate_excel rejects a pinned GPU loopback that sits in a GPU rail subnet.

This check used to say the opposite. A GPU loopback inside the GPU VLAN subnet
was explicitly permitted, on the grounds that the parser generated exactly that
address — `<first GPU VLAN subnet>.4 + core_num`, so `core-01` got
`192.168.0.5`. The permission was accurate about the parser and wrong about the
address: a rail subnet also carries per-switch SVIs from `.2` and, since
ADR-0062, host addresses from `.50`. Three uncoordinated series in one /24, and
at four cores two of them meet — `core-04`'s SVI and `core-01`'s loopback are
both `.5`.

The parser now allocates GPU from the switch's loopback pool
(`VRF_LOOPBACK_BLOCKS['GPU']`), so nothing legitimate lands in a rail subnet
any more and the exemption has no reason to exist. Operators who pinned the old
default need to hear about it, which is what these tests gate.
"""
import sys
from pathlib import Path

import openpyxl

SCRIPTS_DIR = Path(__file__).parent.parent / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

import validate_excel as ve  # noqa: E402

HEADER = ("Switch", "Default", "GPU", "STORAGE")

# Mirrors a collapsed arch: a GPU rail, plus a non-GPU VLAN to prove the
# same-VRF allowance still holds for everything that is not GPU.
VLANS = [
    {"name": "gpu_rail0", "subnet": "192.168.0.0/24", "vrf": "GPU"},
    {"name": "Storage", "subnet": "172.16.180.0/24", "vrf": "STORAGE"},
]


def _ws(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Loopbacks"
    for j, h in enumerate(HEADER, 1):
        ws.cell(1, j, h)
    for i, r in enumerate(rows, 2):
        for j, v in enumerate(r, 1):
            ws.cell(i, j, v)
    return ws


def _errors(rows):
    res = ve.ValidationResult()
    nodes = [{"name": r[0]} for r in rows]
    ve.validate_loopbacks(_ws(rows), nodes, VLANS, {"ns_tiers": 1}, res)
    return [str(e) for e in res.errors if "falls inside VLAN" in str(e)]


def test_the_old_generated_default_is_now_an_error():
    """`192.168.0.5`/`.6` is exactly what the tool used to emit for two cores."""
    errs = _errors([
        ("core-01", "172.16.176.11", "192.168.0.5", "172.16.176.199"),
        ("core-02", "172.16.176.12", "192.168.0.6", "172.16.176.200"),
    ])
    assert len(errs) == 2, errs
    assert "gpu_rail0" in errs[0]
    assert "its own GPU rail subnet" in errs[0], "operator needs to know why"


def test_the_new_default_is_accepted():
    assert _errors([
        ("core-01", "172.16.176.11", "172.16.176.215", "172.16.176.199"),
        ("core-02", "172.16.176.12", "172.16.176.216", "172.16.176.200"),
    ]) == []


def test_any_address_in_the_rail_subnet_is_rejected_not_just_dot_five():
    """The rule is the subnet, not the one address the old formula produced.

    `.50` is the ADR-0062 host start and `.2` is a leaf SVI, so both are
    occupied by something other than a loopback; `.99` is free today and still
    wrong, because the host range grows into it.
    """
    for bad in ("192.168.0.2", "192.168.0.50", "192.168.0.99", "192.168.0.254"):
        errs = _errors([("core-01", "172.16.176.11", bad, "172.16.176.199")])
        assert len(errs) == 1, f"{bad} should be rejected: {errs}"


def test_non_gpu_vrfs_keep_their_same_vrf_allowance():
    """Only GPU lost the exemption. Narrowing it further is a separate call.

    A STORAGE loopback inside the STORAGE VLAN stays legal, so this change
    cannot be blamed for a workbook that starts failing on some other VRF.
    """
    assert _errors([
        ("core-01", "172.16.176.11", "172.16.176.215", "172.16.180.9"),
    ]) == []


def test_a_cross_vrf_overlap_is_still_caught():
    """The original purpose of the check must survive the GPU change."""
    errs = _errors([
        ("core-01", "172.16.176.11", "172.16.176.215", "192.168.0.9"),
    ])
    assert len(errs) == 1
    assert "different VRF" in errs[0]
