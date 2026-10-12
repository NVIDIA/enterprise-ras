# SPDX-FileCopyrightText: Copyright (c) 2026 NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: MIT
"""
The status page is off by default and, when enabled, needs its own real
`status_page_password` (ADR-0091, GitLab #91).

The page is served over plain HTTP and Air forwards it publicly, so:
  * every shipped workbook ships `status_page_enabled = No`;
  * neither the ZTP-server role nor the standalone playbook has a fallback
    value, and both refuse (assert) an unset, placeholder or switch-reused
    password before building .htpasswd — Air and non-Air alike;
  * the password reaches the htpasswd helper on stdin, never in the
    environment (Ansible's shell wrapper carries environment in argv);
  * `check_secrets.py` applies the same rule at generate/deploy time, and the
    development opt-in `ERA_DEV_ALLOW_PLACEHOLDER_SECRETS=1` does not waive it.

The assert tasks are executed with ansible-playbook, not just parsed.
"""
import shutil
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest
import yaml

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "scripts"))

import check_secrets as cs  # noqa: E402

ROLE_TASKS = REPO / "roles" / "ztp-server" / "tasks" / "main.yml"
PLAYBOOK = REPO / "playbooks" / "setup-status-page.yml"
HTPASSWD_SOURCES = [PLAYBOOK, ROLE_TASKS]
ASSERT_NAME = "Refuse a missing or placeholder status_page_password (ADR-0091)"
SWITCH = "Zq9#Synth-Hardened"  # pragma: allowlist secret
DEDICATED = "St@tus-Pag3-Dedicated"  # pragma: allowlist secret

TOOL_WORKBOOKS = sorted(
    list((REPO / "input").glob("*/default/*.xlsx"))
    + list((REPO / "input").glob("*/largescale/*.xlsx"))
    + list((REPO / "input").glob("*.xlsx")))
E2E_FIXTURES = REPO.parent / "release" / "e2e" / "fixtures"
E2E_WORKBOOKS = sorted(E2E_FIXTURES.glob("*.xlsx"))
SHIPPED_WORKBOOKS = TOOL_WORKBOOKS + E2E_WORKBOOKS
GENERATOR = REPO.parent / "data-models" / "generate_arch_excel.py"

needs_ansible = pytest.mark.skipif(not shutil.which("ansible-playbook"),
                                   reason="ansible-playbook required")


def _tasks(path):
    doc = yaml.safe_load(path.read_text())
    return doc if path == ROLE_TASKS else doc[0]["tasks"]


def _task(path, name):
    return next(t for t in _tasks(path) if t.get("name") == name)


def _htpasswd_task(path):
    return next(t for t in _tasks(path) if ".htpasswd" in str(t.get("shell", "")))


def test_every_shipped_workbook_has_the_status_page_off():
    assert len(TOOL_WORKBOOKS) >= 32
    if E2E_FIXTURES.is_dir():
        assert len(E2E_WORKBOOKS) >= 8
    on = []
    for wb_path in SHIPPED_WORKBOOKS:
        ws = openpyxl.load_workbook(wb_path, read_only=True)["Deployment"]
        values = {str(r[0]).strip(): r[1] for r in ws.iter_rows(values_only=True) if r and r[0]}
        if str(values.get("status_page_enabled", "No")).strip().lower() not in ("no", "false"):
            on.append(wb_path.relative_to(REPO.parent))
    assert on == [], f"status page shipped enabled in: {on}"


@pytest.mark.internal_only("data-models/generate_arch_excel.py")
def test_generator_defaults_the_status_page_off():
    src = GENERATOR.read_text()
    assert '("status_page_enabled", "No",' in src


@pytest.mark.parametrize("f", HTPASSWD_SOURCES, ids=lambda p: p.name)
def test_htpasswd_has_no_fallback_and_reads_stdin(f):
    task = _htpasswd_task(f)
    assert "HTPASSWD_PASS" not in str(task.get("environment", {}))
    assert task["args"]["stdin"] == "{{ status_page_password }}"
    assert "sys.stdin" in task["shell"] and task.get("no_log") is True
    text = f.read_text()
    assert "era-status-CHANGE_ME" not in text
    assert "status_page_password | default" not in text


@pytest.mark.parametrize("f", HTPASSWD_SOURCES, ids=lambda p: p.name)
def test_assert_runs_before_htpasswd(f):
    names = [t.get("name") for t in _tasks(f)]
    assert names.index(ASSERT_NAME) < names.index(_htpasswd_task(f)["name"])


@pytest.mark.parametrize("f", HTPASSWD_SOURCES, ids=lambda p: p.name)
def test_assert_placeholder_list_matches_check_secrets(f):
    that = _task(f, ASSERT_NAME)["assert"]["that"]
    listed = next(c for c in that if " not in [" in c)
    assert set(yaml.safe_load(listed.split(" not in ", 1)[1])) == set(cs.STATUS_PAGE_PLACEHOLDERS)


def _run_assert(tmp_path, source, extra_vars):
    task = _task(source, ASSERT_NAME)
    play = [{"hosts": "localhost", "gather_facts": False, "connection": "local",
             "vars": extra_vars, "tasks": [task]}]
    pb = tmp_path / "play.yml"
    pb.write_text(yaml.safe_dump(play, sort_keys=False))
    return subprocess.run(["ansible-playbook", "-i", "localhost,", str(pb)],
                          capture_output=True, text=True, cwd=tmp_path)


CASES = [
    # (label, vars, expected_ok)
    ("disabled-placeholder", {"status_page_enabled": False, "status_page_password": "CHANGE_ME"}, True),  # pragma: allowlist secret
    ("disabled-unset", {"status_page_enabled": False}, True),
    ("enabled-dedicated", {"status_page_enabled": True, "status_page_password": DEDICATED}, True),
    ("enabled-unset", {"status_page_enabled": True}, False),
    ("enabled-empty", {"status_page_enabled": True, "status_page_password": "  "}, False),
    ("enabled-CHANGE_ME", {"status_page_enabled": True, "status_page_password": "CHANGE_ME"}, False),  # pragma: allowlist secret
    ("enabled-old-sentinel", {"status_page_enabled": True, "status_page_password": "era-status-CHANGE_ME"}, False),  # pragma: allowlist secret
    ("enabled-shipped-switch-default", {"status_page_enabled": True, "status_page_password": "Cumu1usLinux!"}, False),  # pragma: allowlist secret
    ("enabled-reuses-switch_password", {"status_page_enabled": True, "status_page_password": SWITCH,
                                        "switch_password": SWITCH}, False),
    ("enabled-reuses-switch_ansible_password", {"status_page_enabled": True, "status_page_password": SWITCH,
                                                "switch_ansible_password": SWITCH}, False),
    ("enabled-multiline-hides-placeholder", {"status_page_enabled": True,
                                             "status_page_password": "password\nA1!suffix"}, False),  # pragma: allowlist secret
    ("enabled-multiline-dedicated", {"status_page_enabled": True, "status_page_password": DEDICATED + "\r\nx"}, False),
    ("enabled-nul", {"status_page_enabled": True, "status_page_password": DEDICATED + "\x00x"}, False),
    ("enabled-equals-default-username", {"status_page_enabled": True, "status_page_password": "era"}, False),  # pragma: allowlist secret
    ("enabled-equals-default-username-case", {"status_page_enabled": True, "status_page_password": "ERA"}, False),  # pragma: allowlist secret
    ("enabled-equals-custom-username", {"status_page_enabled": True, "status_page_username": "ops-admin",
                                        "status_page_password": "ops-admin"}, False),  # pragma: allowlist secret
    ("enabled-equals-custom-username-case", {"status_page_enabled": True, "status_page_username": "Ops-Admin",
                                             "status_page_password": "oPS-ADMIN"}, False),  # pragma: allowlist secret
    ("enabled-dedicated-custom-username", {"status_page_enabled": True, "status_page_username": "ops-admin",
                                           "status_page_password": DEDICATED}, True),
    ("disabled-equals-username", {"status_page_enabled": False, "status_page_password": "era"}, True),  # pragma: allowlist secret
]


@needs_ansible
@pytest.mark.parametrize("deploy_in_air", [True, False], ids=["air", "non-air"])
@pytest.mark.parametrize("source", HTPASSWD_SOURCES, ids=lambda p: p.name)
@pytest.mark.parametrize("label, extra, ok", CASES, ids=[c[0] for c in CASES])
def test_ansible_assert_executes(tmp_path, source, label, extra, ok, deploy_in_air):
    r = _run_assert(tmp_path, source, {**extra, "deploy_in_air": deploy_in_air})
    assert (r.returncode == 0) is ok, r.stdout[-2000:]
    for secret in (DEDICATED, SWITCH):
        assert secret not in r.stdout + r.stderr


def _main_yml(tmp_path, enabled, deploy_in_air, tagged=False):
    d = tmp_path / "output" / "a" / "s" / "inventory" / "group_vars" / "all"
    d.mkdir(parents=True)
    banner = "!unsafe 'x {{ y }}'" if tagged else "'x'"
    (d / "main.yml").write_text(f"deploy_in_air: {str(deploy_in_air).lower()}\n"
                                f"status_page_enabled: {str(enabled).lower()}\n"
                                f"pre_login_message: {banner}\n")
    return tmp_path


@pytest.mark.parametrize("deploy_in_air", [True, False], ids=["air", "non-air"])
@pytest.mark.parametrize("enabled", [True, False])
def test_check_secrets_reads_status_page_enabled(tmp_path, enabled, deploy_in_air):
    root = _main_yml(tmp_path, enabled, deploy_in_air, tagged=True)
    assert cs.status_page_enabled(root, "a", "s") is enabled


def _workbook(tmp_path, value):
    d = tmp_path / "input" / "a" / "s"
    d.mkdir(parents=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Deployment"
    ws.append(["deploy_in_air", "Yes"])
    if value is not None:
        ws.append(["status_page_enabled", value])
    wb.save(d / "a.xlsx")
    return tmp_path


@pytest.mark.parametrize("value,expected", [("Yes", True), ("yes", True), ("No", False), (None, False)])
def test_first_generate_reads_the_workbook(tmp_path, value, expected):
    root = _workbook(tmp_path, value)
    assert cs.status_page_enabled(root, "a", "s") is expected


def test_workbook_enabling_beats_a_stale_disabled_inventory(tmp_path):
    root = _workbook(_main_yml(tmp_path, False, True), "Yes")
    assert cs.status_page_enabled(root, "a", "s") is True


def test_nothing_generated_and_no_workbook_is_unknown(tmp_path):
    assert cs.status_page_enabled(tmp_path, "a", "s") is None


def test_generate_step_0b_refuses_a_workbook_that_newly_enables_the_page(tmp_path, monkeypatch, capsys):
    monkeypatch.delenv(cs.DEV_OPT_IN_ENV, raising=False)
    root = _workbook(_main_yml(tmp_path, False, True), "Yes")
    secrets = root / "output" / "a" / "s" / "inventory" / "group_vars" / "all" / "secrets.yml"
    secrets.write_text(yaml.safe_dump({"switch_password": SWITCH, "switch_ansible_password": SWITCH}))
    monkeypatch.setattr(cs, "__file__", str(root / "scripts" / "check_secrets.py"))
    assert cs.main(["--arch", "a", "--site", "s"]) == 1
    assert "status_page_password is not set" in capsys.readouterr().out


@pytest.mark.parametrize("secrets, fragment", [
    ({}, "status_page_password is not set"),
    ({"status_page_password": ""}, "status_page_password is empty"),
    ({"status_page_password": "CHANGE_ME"}, cs.PLACEHOLDER_PROBLEM),  # pragma: allowlist secret
    ({"status_page_password": "Cumu1usLinux!"}, cs.PLACEHOLDER_PROBLEM),  # pragma: allowlist secret
    ({"status_page_password": "era-status-CHANGE_ME"}, cs.PLACEHOLDER_PROBLEM),  # pragma: allowlist secret
    ({"status_page_password": SWITCH, "switch_password": SWITCH}, "reuses switch_password"),
    ({"status_page_password": "era"}, "equals the username"),  # pragma: allowlist secret
    ({"status_page_password": "ERA"}, "equals the username"),  # pragma: allowlist secret
    ({"status_page_username": "Ops-Admin", "status_page_password": "oPS-ADMIN"}, "equals the username"),  # pragma: allowlist secret
    ({"status_page_password": "password\nA1!suffix"}, "line break"),  # pragma: allowlist secret
    ({"status_page_password": DEDICATED + "\nx"}, "line break"),
])
def test_status_page_password_findings(secrets, fragment):
    findings, unchecked = cs.check_status_page_password(secrets)
    assert unchecked == [] and len(findings) == 1 and fragment in findings[0], findings


def test_dedicated_status_page_password_passes():
    assert cs.check_status_page_password({"status_page_password": DEDICATED,
                                          "switch_password": SWITCH}) == ([], [])


def _cli(*args, opt_in=False):
    import os
    env = {k: v for k, v in os.environ.items() if k != cs.DEV_OPT_IN_ENV}
    if opt_in:
        env[cs.DEV_OPT_IN_ENV] = "1"
    return subprocess.run([sys.executable, str(REPO / "scripts" / "check_secrets.py"), *args],
                          capture_output=True, text=True, env=env)


@pytest.mark.parametrize("opt_in", [False, True], ids=["no-opt-in", "dev-opt-in"])
def test_enabled_status_page_placeholder_is_never_waived(tmp_path, opt_in):
    p = tmp_path / "secrets.yml"
    p.write_text(f'switch_password: "{SWITCH}"\nswitch_ansible_password: "{SWITCH}"\n'
                 'status_page_password: "CHANGE_ME"\n')  # pragma: allowlist secret
    r = _cli("--file", str(p), "--status-page-enabled", opt_in=opt_in)
    assert r.returncode == 1 and "status_page_password" in r.stdout
    assert "CHANGE_ME" not in r.stdout


def test_disabled_status_page_placeholder_is_not_gated(tmp_path):
    p = tmp_path / "secrets.yml"
    p.write_text(f'switch_password: "{SWITCH}"\nswitch_ansible_password: "{SWITCH}"\n'
                 'status_page_password: "CHANGE_ME"\n')  # pragma: allowlist secret
    assert _cli("--file", str(p)).returncode == 0


def test_status_page_password_seeded_in_secrets():
    """status_page_password must be present in every source secrets.yml."""
    for sec in [REPO / "inventories" / "secrets.yml", REPO / "inventories" / "secrets.yml.example"]:
        assert "status_page_password" in sec.read_text(), f"{sec} lacks status_page_password"


@needs_ansible
@pytest.mark.skipif(not shutil.which("ansible-vault"), reason="ansible-vault required")
@pytest.mark.parametrize("source", HTPASSWD_SOURCES, ids=lambda p: p.name)
@pytest.mark.parametrize("password, ok", [("era", False), ("Era", False), (DEDICATED, True)],
                         ids=["equals-username", "equals-username-case", "dedicated"])
def test_vaulted_secrets_defer_to_the_runtime_assert(tmp_path, source, password, ok):
    """A whole-file vault is skipped by the preflight; the deploy assert must still check it."""
    secrets = tmp_path / "secrets.yml"
    secrets.write_text(yaml.safe_dump({"status_page_username": "era", "status_page_password": password,
                                       "switch_password": SWITCH}))
    vault_pass = tmp_path / "vault-pass"
    vault_pass.write_text("synthetic-vault-pass\n")  # pragma: allowlist secret
    subprocess.run(["ansible-vault", "encrypt", "--vault-password-file", str(vault_pass), str(secrets)],
                   check=True, capture_output=True, cwd=tmp_path)
    assert cs.is_vault_encrypted(secrets)
    findings, note = cs.check_secrets_file(secrets, status_page=True)
    assert findings == [] and "vault-encrypted" in note

    play = [{"hosts": "localhost", "gather_facts": False, "connection": "local",
             "vars": {"status_page_enabled": True}, "vars_files": [str(secrets)],
             "tasks": [_task(source, ASSERT_NAME)]}]
    pb = tmp_path / "play.yml"
    pb.write_text(yaml.safe_dump(play, sort_keys=False))
    r = subprocess.run(["ansible-playbook", "-i", "localhost,", "--vault-password-file", str(vault_pass), str(pb)],
                       capture_output=True, text=True, cwd=tmp_path)
    assert (r.returncode == 0) is ok, r.stdout[-2000:]
    if not ok:
        assert "status_page_username" in r.stdout
