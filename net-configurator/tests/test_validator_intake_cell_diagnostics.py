# SPDX-FileCopyrightText: Copyright (c) 2026 NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: MIT
"""GitLab #97 / #100 (H9, H12): the validator names intake sheets and cells.

Each case starts from a shipped intake workbook that passes
`make validate-excel`, changes one cell, and asserts an error that names the
intake sheet and cell the operator has to fix. Before this, the findings
named the in-memory legacy views (`VLANs & Profiles`, `Nodes`, `Wire Map`)
and their row numbers, and several malformed values were skipped without a
word (`continue`) or ended a table silently.
"""

import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

REPO = Path(__file__).resolve().parent.parent
SCRIPTS = REPO / "scripts"
for _path in (SCRIPTS, REPO / "tests"):
    if str(_path) not in sys.path:
        sys.path.insert(0, str(_path))

import workbook_access  # noqa: E402
from intake_reader import column_map, find_section, header_row  # noqa: E402
from intake_to_legacy import Origins  # noqa: E402
from validate_excel import (  # noqa: E402
    ValidationResult, _add_intake_legacy_views, validate_excel, validate_nodes)

BASE = REPO / "input" / "2-4-3-200" / "default" / "2-4-3-200.xlsx"


def _network_row(ws, vlan):
    col = column_map(ws)["vlan"]
    for r in range(header_row(ws) + 1, ws.max_row + 1):
        if ws.cell(r, col).value == vlan:
            return r
    raise AssertionError(f"VLAN {vlan} not in Network Data")


def _section_row(ws, title, first_cell):
    """Row of the `first_cell` data row in the stacked section `title`."""
    start = find_section(ws, title)
    for r in range(start + 2, ws.max_row + 1):
        if ws.cell(r, 1).value == first_cell:
            return r
    raise AssertionError(f"{first_cell} not in {title}")


def _col(ws, header, header_at=None):
    row = header_at or header_row(ws)
    for c in range(1, ws.max_column + 1):
        if str(ws.cell(row, c).value or "").strip().lower() == header.lower():
            return c
    raise AssertionError(f"{header} not in row {row}")


def _letter(col):
    return openpyxl.utils.get_column_letter(col)


def _errors(tmp_path, edit):
    wb = openpyxl.load_workbook(BASE)
    edit(wb)
    out = tmp_path / "mutated.xlsx"
    wb.save(out)
    return validate_excel(str(out)).errors


def _has(errors, *parts):
    return [e for e in errors if all(p in e for p in parts)]


def test_base_workbook_passes_make_validate_excel(tmp_path):
    copy = tmp_path / "base.xlsx"
    openpyxl.load_workbook(BASE).save(copy)
    proc = subprocess.run(["make", "--no-print-directory", "validate-excel",
                           f"EXCEL={copy}"], cwd=REPO, capture_output=True,
                          text=True, timeout=300)
    assert proc.returncode == 0, proc.stdout + proc.stderr
    assert "All checks passed" in proc.stdout


def test_nonnumeric_vni(tmp_path):
    def edit(wb):
        ws = wb["Network Data"]
        ws.cell(_network_row(ws, 300), column_map(ws)["vni(d)"]).value = "abc"
    errors = _errors(tmp_path, edit)
    ws = openpyxl.load_workbook(BASE)["Network Data"]
    ref = f"'Network Data'!{_letter(column_map(ws)['vni(d)'])}{_network_row(ws, 300)}"
    assert _has(errors, "[Network Data]", ref, "VNI 'abc' on VLAN 300"), errors


@pytest.mark.parametrize("value, phrase", [
    ("abc", "is not a whole VLAN number"),
    ("300,400", "lists more than one VLAN"),
    ("  ", "holds only whitespace"),
])
def test_untagged_vlan_must_be_one_whole_vlan(tmp_path, value, phrase):
    def edit(wb):
        ws = wb["Network Data"]
        row = _section_row(ws, "Port Profiles", "CPU/In-Band Network")
        hdr = find_section(ws, "Port Profiles") + 1
        ws.cell(row, _col(ws, "Untagged VLAN", hdr)).value = value
    errors = _errors(tmp_path, edit)
    ws = openpyxl.load_workbook(BASE)["Network Data"]
    hdr = find_section(ws, "Port Profiles") + 1
    ref = (f"'Network Data'!{_letter(_col(ws, 'Untagged VLAN', hdr))}"
           f"{_section_row(ws, 'Port Profiles', 'CPU/In-Band Network')}")
    assert _has(errors, "[Network Data]", ref, "Untagged VLAN", phrase), errors


@pytest.mark.parametrize("value, phrase", [
    ("abc", "is not a number"),
    (900.5, "is a decimal number"),
    (None, "VLAN is blank"),
    (0, "VLAN ID 0 out of valid range"),
    (4095, "VLAN ID 4095 out of valid range"),
    (-5, "VLAN ID -5 out of valid range"),
])
def test_malformed_vlan_id(tmp_path, value, phrase):
    def edit(wb):
        ws = wb["Network Data"]
        ws.cell(_network_row(ws, 400), column_map(ws)["vlan"]).value = value
    errors = _errors(tmp_path, edit)
    ws = openpyxl.load_workbook(BASE)["Network Data"]
    ref = f"'Network Data'!{_letter(column_map(ws)['vlan'])}{_network_row(ws, 400)}"
    assert _has(errors, "[Network Data]", ref, phrase), errors


def test_blank_row_inside_vlan_table(tmp_path):
    def edit(wb):
        ws = wb["Network Data"]
        row = _network_row(ws, 400)
        for c in range(1, ws.max_column + 1):
            ws.cell(row, c).value = None
    errors = _errors(tmp_path, edit)
    row = _network_row(openpyxl.load_workbook(BASE)["Network Data"], 400)
    assert _has(errors, "[Network Data]", f"Row {row + 1} has VLAN data below "
                f"the blank row {row}"), errors


def _vrf_cell(ws, vrf, header):
    hdr = find_section(ws, "VRFs") + 1
    return _section_row(ws, "VRFs", vrf), _col(ws, header, hdr)


def test_vrf_vlan_equal_to_tenant_vlan(tmp_path):
    def edit(wb):
        ws = wb["Network Data"]
        ws.cell(*_vrf_cell(ws, "INBAND", "VLAN")).value = 300
    errors = _errors(tmp_path, edit)
    ws = openpyxl.load_workbook(BASE)["Network Data"]
    row, col = _vrf_cell(ws, "INBAND", "VLAN")
    assert _has(errors, "[Network Data]", f"'Network Data'!{_letter(col)}{row}",
                "VLAN 300 of VRF 'INBAND' is also tenant VLAN"), errors


def test_two_vrfs_share_an_l3_vni_vlan(tmp_path):
    def edit(wb):
        ws = wb["Network Data"]
        ws.cell(*_vrf_cell(ws, "GPU", "VLAN")).value = 3002
    errors = _errors(tmp_path, edit)
    ws = openpyxl.load_workbook(BASE)["Network Data"]
    row, col = _vrf_cell(ws, "GPU", "VLAN")
    assert _has(errors, "[Network Data]", f"'Network Data'!{_letter(col)}{row}",
                "VLAN 3002 of VRF 'GPU' is also the VLAN of VRF 'INBAND'"), errors


def test_l2_vni_equal_to_l3_vni(tmp_path):
    def edit(wb):
        ws = wb["Network Data"]
        ws.cell(_network_row(ws, 300), column_map(ws)["vni(d)"]).value = 5002
    errors = _errors(tmp_path, edit)
    ws = openpyxl.load_workbook(BASE)["Network Data"]
    row, col = _vrf_cell(ws, "INBAND", "L3 VNI")
    assert _has(errors, "[Network Data]", f"'Network Data'!{_letter(col)}{row}",
                "L3 VNI 5002 of VRF 'INBAND' is also the VNI of VLAN 300"), errors


def test_enabled_outside_documented_values(tmp_path):
    def edit(wb):
        ws = wb["Switch Data"]
        ws.cell(header_row(ws) + 1, column_map(ws)["enabled"]).value = "Y"
    errors = _errors(tmp_path, edit)
    ws = openpyxl.load_workbook(BASE)["Switch Data"]
    ref = f"'Switch Data'!{_letter(column_map(ws)['enabled'])}{header_row(ws) + 1}"
    assert _has(errors, "[Switch Data]", ref, "Enabled 'Y' for core-01"), errors


def test_blank_display_in_air(tmp_path):
    def edit(wb):
        ws = wb["Wiremap"]
        ws.cell(header_row(ws) + 1, column_map(ws)["display in air"]).value = None
    errors = _errors(tmp_path, edit)
    ws = openpyxl.load_workbook(BASE)["Wiremap"]
    ref = f"Wiremap!{_letter(column_map(ws)['display in air'])}{header_row(ws) + 1}"
    assert _has(errors, "[Wiremap]", ref, "Display in Air is blank"), errors


def test_existing_row_findings_name_intake_rows(tmp_path):
    """A pre-existing check: its label and rows are the intake sheet's."""
    def edit(wb):
        ws = wb["Network Data"]
        name = column_map(ws)["fabric name"]
        ws.cell(_network_row(ws, 400), name).value = "CPU/In-Band"
    errors = _errors(tmp_path, edit)
    ws = openpyxl.load_workbook(BASE)["Network Data"]
    rows = sorted([_network_row(ws, 300), _network_row(ws, 400)])
    assert _has(errors, "[Network Data]", "Duplicate VLAN name 'cpu/in-band'",
                f"rows: [{rows[0]}, {rows[1]}]"), errors
    assert not [e for e in errors if "VLANs & Profiles" in e], errors


def test_blank_display_in_air_is_not_required_on_a_physical_site(tmp_path):
    """`Display in Air` feeds only the Air topology (ADR-0093 physical sites)."""
    def edit(wb):
        workbook_access.set_setting(wb, "deploy_in_air", "No")
        ws = wb["Wiremap"]
        ws.cell(header_row(ws) + 1, column_map(ws)["display in air"]).value = None
    errors = _errors(tmp_path, edit)
    assert not _has(errors, "Display in Air is blank"), errors


@pytest.mark.parametrize("value, enabled", [
    (False, False), (0, False), ("No", False),
    (True, True), (1, True), ("Yes", True), (None, True),
])
def test_enabled_is_read_as_the_generator_reads_it(value, enabled):
    """Boolean False and numeric 0 disable the node in excel_parser too."""
    wb = openpyxl.load_workbook(BASE)
    ws = wb["Switch Data"]
    ws.cell(header_row(ws) + 1, column_map(ws)["enabled"]).value = value
    result = ValidationResult()
    result.origins = _add_intake_legacy_views(wb)
    nodes = validate_nodes(wb["Nodes"], result, settings=None)
    core = next(n for n in nodes if n["name"] == "core-01")
    assert core["enabled"] is enabled
    assert not _has(result.errors, "Enabled"), result.errors


def test_switch_hardware_finding_names_the_wiremap_cell(tmp_path):
    def edit(wb):
        ws = wb["Wiremap"]
        ws.cell(header_row(ws) + 1, column_map(ws)["port (z)"]).value = "swp70"
    errors = _errors(tmp_path, edit)
    ws = openpyxl.load_workbook(BASE)["Wiremap"]
    ref = f"Wiremap!{_letter(column_map(ws)['port (z)'])}{header_row(ws) + 1}"
    assert _has(errors, "[Switch hardware]", ref, "uses swp70"), errors
    assert not _has(errors, "Wire Map row"), errors


def test_mgmt_ip_outside_oob_finding_names_the_intake_cell(tmp_path):
    def edit(wb):
        ws = wb["Switch Data"]
        ws.cell(header_row(ws) + 1, column_map(ws)["mgmt ip"]).value = "10.250.250.250"
    errors = _errors(tmp_path, edit)
    ws = openpyxl.load_workbook(BASE)["Switch Data"]
    ref = f"'Switch Data'!{_letter(column_map(ws)['mgmt ip'])}{header_row(ws) + 1}"
    assert _has(errors, "[Cross-sheet]", ref, "10.250.250.250",
                "is not within any OOB VLAN subnet"), errors


def test_row_lists_keep_each_separator():
    origins = Origins()
    for legacy, intake in ((3, 13), (4, 14), (5, 15)):
        origins.add("VLANs & Profiles", legacy, "Network Data", intake)
    result = ValidationResult()
    result.origins = origins
    result.error("VLANs & Profiles", "VLAN 7 repeats on rows 3, 4 and 5")
    assert result.errors == ["[Network Data] VLAN 7 repeats on rows 13, 14 and 15"]


def test_ambiguous_network_data_header_is_a_workbook_finding(tmp_path):
    """`header_row` refuses a row 1 that mixes group labels and column names."""
    def edit(wb):
        wb["Network Data"]["B1"] = "L2"
    errors = _errors(tmp_path, edit)
    assert _has(errors, "[Workbook] intake-form workbook could not be read",
                "Sheet 'Network Data': row 1 mixes known group labels"), errors


def _mgmt_ip_ref(wb, sheet, row):
    ws = wb[sheet]
    key = "mgmt ip" if sheet == "Switch Data" else "ip"
    return f"'{sheet}'!{_letter(column_map(ws)[key])}{row}"


def test_duplicate_mgmt_ip_names_both_switch_data_cells(tmp_path):
    def edit(wb):
        ws = wb["Switch Data"]
        col, first = column_map(ws)["mgmt ip"], header_row(ws) + 1
        ws.cell(first + 1, col).value = ws.cell(first, col).value
    errors = _errors(tmp_path, edit)
    wb = openpyxl.load_workbook(BASE)
    first = header_row(wb["Switch Data"]) + 1
    a, b = (_mgmt_ip_ref(wb, "Switch Data", r) for r in (first, first + 1))
    assert _has(errors, "Duplicate management IP", f"used by: {a} (core), {b} (core)"), errors


def test_duplicate_mgmt_ip_across_sheets_names_each_sheet(tmp_path):
    wb = openpyxl.load_workbook(BASE)
    hosts = wb["Host Data"]
    ip_col = column_map(hosts)["ip"]
    host_row = next(r for r in range(header_row(hosts) + 1, hosts.max_row + 1)
                    if hosts.cell(r, ip_col).value)
    function = hosts.cell(host_row, column_map(hosts)["function"]).value

    def edit(wb):
        ws = wb["Switch Data"]
        wb["Host Data"].cell(host_row, ip_col).value = ws.cell(
            header_row(ws) + 1, column_map(ws)["mgmt ip"]).value
    errors = _errors(tmp_path, edit)
    a = _mgmt_ip_ref(wb, "Switch Data", header_row(wb["Switch Data"]) + 1)
    b = _mgmt_ip_ref(wb, "Host Data", host_row)
    assert _has(errors, "Duplicate management IP",
                f"used by: {a} (core), {b} ({function})"), errors
