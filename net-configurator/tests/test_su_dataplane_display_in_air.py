# SPDX-FileCopyrightText: Copyright (c) 2026 NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: MIT
"""Warn per SU when every GPU/CPU link is `Display in Air = No`, offer to fix (ERA-130).

An SU whose compute nodes are cabled to the GPU rails and the CPU/in-band
leaves but has all of those rows marked No builds an Air sim with OOB-only
compute nodes. That is almost always a copy-paste slip, so the validator warns
and — interactively, or with AIR_DISPLAY_FIX=yes — flips the rows to Yes.
Non-interactive runs (CI) only warn.
"""
import io
import shutil
import sys
from pathlib import Path

import openpyxl
import pytest

NC = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(NC / "scripts"))

import validate_excel as ve  # noqa: E402

HEADER = ["Display in Air", "System Name (A)", "Port (A)", "Port Side (A)",
          "Cable Split (A)", "System Name (B)", "Port (B)", "Port Side (B)",
          "Cable Split (B)", "Network Profile"]


def _sheet(rows):
    """rows: (display, node, port, profile). Real openpyxl sheet — the column map reads it."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Wire Map"
    ws.append(HEADER)
    for disp, node, port, profile in rows:
        ws.append([disp, node, port, "", "", "leaf-01", "swp1", "", "", profile])
    return ws


def _nodes(*names, function="GPU"):
    return [{"name": n, "function": function, "enabled": True} for n in names]


def _run(rows, nodes, settings=None):
    res = ve.ValidationResult()
    ve.validate_su_dataplane_display_in_air(_sheet(rows), nodes, res, settings)
    return res


ALL_NO = [
    ("No", "su-01-node-01", "eth1", "GPU Network"),
    ("No", "su-01-node-01", "eth2", "CPU/In-Band Network"),
    ("No", "su-01-node-02", "eth1", "GPU Network"),
    ("No", "su-01-node-02", "eth2", "CPU/In-Band Network"),
    ("No", "su-01-node-01", "eth0", "OOB / IPMI"),
]
TWO_NODES = _nodes("su-01-node-01", "su-01-node-02")


def test_all_no_warns_once_per_su_and_tells_operator_to_set_yes():
    res = _run(ALL_NO, TWO_NODES)
    assert len(res.warnings) == 1
    msg = res.warnings[0]
    assert "su-01" in msg and "4 gpu+cpu links" in msg and "2 compute node(s)" in msg
    assert "Display in Air = Yes" in msg and "AIR_DISPLAY_FIX=yes" in msg
    assert res.errors == []  # a warning, never a hard stop
    assert sorted(res.air_display_fix_rows) == [
        ("su-01-node-01", "eth1"), ("su-01-node-01", "eth2"),
        ("su-01-node-02", "eth1"), ("su-01-node-02", "eth2")]


def test_oob_rows_are_not_offered_for_fixing():
    res = _run(ALL_NO, TWO_NODES)
    assert ("su-01-node-01", "eth0") not in res.air_display_fix_rows


def test_one_yes_anywhere_in_the_su_is_enough():
    rows = [("Yes" if i == 2 else d, n, p, pr) for i, (d, n, p, pr) in enumerate(ALL_NO)]
    res = _run(rows, TWO_NODES)
    assert res.warnings == [] and res.air_display_fix_rows == []


@pytest.mark.parametrize("value", ["yes", "YES", " Yes "])
def test_yes_is_case_insensitive_like_the_parser(value):
    rows = [(value, "su-01-node-01", "eth1", "GPU Network"),
            ("No", "su-01-node-01", "eth2", "CPU/In-Band Network")]
    assert _run(rows, _nodes("su-01-node-01")).warnings == []


def test_sus_are_judged_independently():
    rows = ALL_NO + [("Yes", "su-02-node-01", "eth1", "GPU Network"),
                     ("No", "su-02-node-01", "eth2", "CPU/In-Band Network")]
    res = _run(rows, TWO_NODES + _nodes("su-02-node-01"))
    assert len(res.warnings) == 1 and "su-01" in res.warnings[0]
    assert all(n.startswith("su-01") for n, _ in res.air_display_fix_rows)


def test_gpu_only_su_still_warns_with_gpu_label():
    rows = [("No", "su-01-node-01", "eth1", "GPU Network"),
            ("No", "su-01-node-01", "eth3", "GPU Network")]
    res = _run(rows, _nodes("su-01-node-01"))
    assert len(res.warnings) == 1 and "2 gpu links" in res.warnings[0]


def test_non_compute_nodes_are_ignored():
    rows = [("No", "k8s-master-01", "eth1", "CPU/In-Band Network"),
            ("No", "storage-01", "eth1", "Storage")]
    res = _run(rows, _nodes("k8s-master-01", function="K8s") + _nodes("storage-01", function="Storage"))
    assert res.warnings == []


def test_disabled_nodes_are_ignored():
    nodes = _nodes("su-01-node-01", "su-01-node-02")
    nodes[0]["enabled"] = False
    nodes[1]["enabled"] = False
    assert _run(ALL_NO, nodes).warnings == []


@pytest.mark.parametrize("value", ["No", "false", "0"])
def test_silent_when_not_deploying_in_air(value):
    assert _run(ALL_NO, TWO_NODES, {"deploy_in_air": value}).warnings == []


def test_absent_deploy_key_defaults_to_air_like_the_oob_rule():
    assert len(_run(ALL_NO, TWO_NODES, {}).warnings) == 1


def test_node_without_su_prefix_is_its_own_group():
    rows = [("No", "gpu-node-07", "eth1", "GPU Network"),
            ("No", "gpu-node-07", "eth2", "CPU/In-Band Network")]
    res = _run(rows, _nodes("gpu-node-07"))
    assert len(res.warnings) == 1 and "'gpu-node-07'" in res.warnings[0]


# --- override decision -------------------------------------------------------

class _Tty(io.StringIO):
    def isatty(self):
        return True


@pytest.mark.parametrize("value", ["yes", "YES", "y", "true", "1"])
def test_env_yes_decides_without_prompt(value):
    assert ve.decide_air_display_fix(4, env={"AIR_DISPLAY_FIX": value}, stdin=io.StringIO()) == "yes"


@pytest.mark.parametrize("value", ["no", "n", "false", "0", "maybe"])
def test_env_no_or_garbage_declines(value, capsys):
    assert ve.decide_air_display_fix(4, env={"AIR_DISPLAY_FIX": value}, stdin=_Tty()) == "no"


def test_non_tty_never_prompts_and_declines(monkeypatch):
    monkeypatch.setattr("builtins.input", lambda *_: pytest.fail("prompted in non-interactive run"))
    assert ve.decide_air_display_fix(4, env={}, stdin=io.StringIO()) == "no"


@pytest.mark.parametrize("answer, expected", [("y", "yes"), ("yes", "yes"), ("", "no"), ("n", "no")])
def test_tty_prompt_mentions_yes_and_default_is_no(monkeypatch, answer, expected):
    seen = {}

    def fake_input(prompt):
        seen["prompt"] = prompt
        return answer
    monkeypatch.setattr("builtins.input", fake_input)
    assert ve.decide_air_display_fix(4, env={}, stdin=_Tty()) == expected
    assert "Display in Air = Yes" in seen["prompt"] and "[y/N]" in seen["prompt"]


# --- applying the fix to a real intake workbook -------------------------------

SAMPLE = NC / "input" / "2-8-9-800" / "default" / "2-8-9-800.xlsx"


def _flip_su01_to_no(path):
    import intake_to_legacy
    wb = openpyxl.load_workbook(path)
    ws = wb["Wiremap"]
    cols = intake_to_legacy.column_map(ws)
    flipped = 0
    for r in range(intake_to_legacy.header_row(ws) + 1, ws.max_row + 1):
        name = str(ws.cell(r, cols["system name (a)"]).value or "")
        prof = str(ws.cell(r, cols["network profile"]).value or "").lower()
        if name.startswith("su-01-") and ("gpu" in prof or "cpu" in prof):
            ws.cell(r, cols["display in air"]).value = "No"
            flipped += 1
    wb.save(path)
    return flipped


@pytest.mark.skipif(not SAMPLE.exists(), reason="sample workbook not present")
def test_shipped_workbooks_do_not_trigger_the_warning():
    for xlsx in sorted((NC / "input").glob("*/default/*.xlsx")):
        res = ve.validate_excel(xlsx)
        assert not [w for w in res.warnings if "Display in Air = No" in w], xlsx.name


@pytest.mark.skipif(not SAMPLE.exists(), reason="sample workbook not present")
def test_end_to_end_warn_then_fix_on_intake_workbook(tmp_path):
    work = tmp_path / "t.xlsx"
    shutil.copy(SAMPLE, work)
    flipped = _flip_su01_to_no(work)
    assert flipped > 0

    res = ve.validate_excel(work)
    hits = [w for w in res.warnings if "Display in Air = No" in w]
    assert len(hits) == 1 and "'su-01'" in hits[0]
    assert len(res.air_display_fix_rows) == flipped

    changed = ve.apply_air_display_fix(work, res.air_display_fix_rows)
    assert changed == flipped

    after = ve.validate_excel(work)
    assert not [w for w in after.warnings if "Display in Air = No" in w]
    assert after.air_display_fix_rows == []
    # Backup landed under net-configurator/archive/<date>/ before the write.
    backups = list((NC / "archive").glob(f"*/{work.name}.*.bak"))
    assert backups, "workbook was rewritten without an archive/ backup"
    for b in backups:
        b.unlink()
