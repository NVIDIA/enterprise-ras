# SPDX-FileCopyrightText: Copyright (c) 2026 NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: MIT
"""The Deployment sheet carries the settings that cannot be derived.

Instruction rows are `#`-prefixed, matching the convention Custom_Config
already uses, and must not be read as settings.
"""
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).parent.parent / "scripts"))

from intake_reader import read_deployment


def _sheet():
    wb = openpyxl.Workbook()
    ws = wb.create_sheet("Deployment")
    ws.append(["# ── OPTIONAL SHEETS ──"])
    ws.append(["#   Custom_Config    nv commands appended last"])
    ws.append(["Setting", "Value"])
    ws.append(["timezone", "Etc/Zulu"])
    ws.append(["ldap_enabled", "No"])
    ws.append([None, None])
    ws.append(["core", "5.18.0"])
    return ws


def test_settings_are_read_as_name_value_pairs():
    assert read_deployment(_sheet())["timezone"] == "Etc/Zulu"


def test_instruction_rows_are_skipped():
    values = read_deployment(_sheet())
    assert not any(k.startswith("#") for k in values)


def test_blank_row_does_not_terminate_the_table():
    assert read_deployment(_sheet())["core"] == "5.18.0"


def test_a_setting_named_setting_is_read_not_skipped():
    """A legitimate setting named 'Setting' is read when its value is not 'Value'."""
    wb = openpyxl.Workbook()
    ws = wb.create_sheet("Deployment")
    ws.append(["Setting", "Value"])  # header row
    ws.append(["Setting", "some_config"])  # a setting named "Setting"
    result = read_deployment(ws)
    assert "Setting" in result
    assert result["Setting"] == "some_config"


def test_duplicate_keys_raise():
    """A duplicate key raises, naming both the key and the conflicting values."""
    wb = openpyxl.Workbook()
    ws = wb.create_sheet("Deployment")
    ws.append(["Setting", "Value"])
    ws.append(["timezone", "Etc/Zulu"])
    ws.append(["timezone", "America/New_York"])  # duplicate
    try:
        read_deployment(ws)
        assert False, "Expected DuplicateKeyError"
    except ValueError as e:
        assert "timezone" in str(e)
        assert "Etc/Zulu" in str(e) or "America/New_York" in str(e)


def test_section_banners_are_skipped_but_blank_valued_settings_survive():
    """The banner rule needs BOTH conditions; each alone loses real data.

    `netq_ip` ships with an empty value in all 14 shipped workbooks, so
    skipping every value-less row would silently drop it -- and it is one of
    the settings the converter was already destroying before this branch.
    A setting legitimately named "Setting" must also survive (see the test
    above), so skipping by name alone is wrong too.
    """
    wb = openpyxl.Workbook()
    ws = wb.create_sheet("Deployment")
    ws.append(["Setting", "Value"])
    ws.append(["GENERAL", None])          # banner: known name, no value
    ws.append(["site_name", "default"])
    ws.append(["TELEMETRY", None])        # banner
    ws.append(["netq_ip", None])          # real setting, legitimately blank
    result = read_deployment(ws)

    assert "GENERAL" not in result and "TELEMETRY" not in result
    assert "netq_ip" in result, (
        "a blank-valued setting was dropped -- banners must be matched by "
        "name, not by having no value")
    assert result["netq_ip"] is None
    assert result["site_name"] == "default"


def test_numeric_zero_in_enabled_means_disabled_not_active():
    """`Enabled = 0` must be Disabled.

    `str(value or 'Yes')` relies on falsiness, and 0 is falsy -- so a node
    explicitly disabled with a numeric 0 came back Active. openpyxl returns
    real ints for numeric cells. (Greptile P1)
    """
    import sys
    from pathlib import Path
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
    from excel_parser import _read_intake_node_rows

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Function", "Name", "Enabled"])
    ws.append(["gpu", "node-yes", "Yes"])
    ws.append(["gpu", "node-zero", 0])
    ws.append(["gpu", "node-blank", None])

    field_map = {"role": "function", "name": "name", "enabled": "enabled",
                 "mac_address": "mac address", "mgmt_ip": "ip",
                 "prefix": "prefix", "gateway": "gateway", "ztp": "ztp",
                 "oob_vlan": "oob vlan"}
    got = {n["name"]: n["status"] for n in _read_intake_node_rows(ws, field_map)}

    assert got["node-zero"] == "Disabled", (
        f"numeric 0 was read as {got['node-zero']!r}, not Disabled")
    assert got["node-yes"] == "Active"
    assert got["node-blank"] == "Active", "blank should still default to Active"
