# SPDX-FileCopyrightText: Copyright (c) 2026 NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: MIT
"""Copy the downloadable sheets into an intake workbook and render their effects."""
import os
import subprocess
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
SHEETS = ('Prefix lists', 'Route policy', 'Community lists', 'ACLs', 'Custom_Config')


@pytest.fixture(scope='module')
def rendered_sample(tmp_path_factory):
    tmp = tmp_path_factory.mktemp('optional-sample')
    source = openpyxl.load_workbook(ROOT / 'docs/optional-overrides.xlsx')
    workbook = openpyxl.load_workbook(ROOT / 'input/2-4-3-200/default/2-4-3-200.xlsx')
    try:
        for name in SHEETS:
            sheet = workbook.create_sheet(name)
            for row in source[name].iter_rows(values_only=True):
                sheet.append(row)
        path = tmp / 'example.xlsx'
        workbook.save(path)
    finally:
        source.close()
        workbook.close()
    result = subprocess.run(['make', 'validate-excel', f'EXCEL={path}'],
                            cwd=ROOT, capture_output=True, text=True)
    assert result.returncode == 0, result.stdout + result.stderr
    with pytest.MonkeyPatch.context() as patch:
        patch.syspath_prepend(str(ROOT / 'scripts'))
        from excel_parser import process_excel_template
    inventory = tmp / 'inventory'
    process_excel_template(path, inventory, site='optional-example')
    configs = tmp / 'configs'
    env = dict(os.environ, ANSIBLE_LOCAL_TEMP=str(tmp / 'ansible-local'))
    result = subprocess.run(
        ['ansible-playbook', 'playbooks/generate-cli-configs.yml',
         '-i', str(inventory / 'hosts'), '-e', f'config_output_dir={configs}'],
        cwd=ROOT, env=env, capture_output=True, text=True,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    return {p.name: p.read_text() for p in configs.glob('*-config.sh')}


@pytest.mark.parametrize('fragment', [
    'router policy prefix-list EXAMPLE_PREFIX rule 10',
    'router policy route-map EXAMPLE_POLICY rule 10',
    'router policy community-list EXAMPLE_COMMUNITY rule 10',
    'acl example-ssh-allow rule 10',
    "system message pre-login 'Example lab - authorized users only'",
])
def test_downloadable_sheet_examples_reach_switch_config(rendered_sample, fragment):
    assert fragment in rendered_sample['core-01-config.sh']


def test_sample_preserves_default_rules_and_does_not_bind_example_route_map(rendered_sample):
    core = rendered_sample['core-01-config.sh']
    assert 'router policy prefix-list ERA_PREFIXES rule ' in core
    assert 'acl acl-default-whitelist' in core
    assert 'acl acl-default-dos' in core
    assert not any('EXAMPLE_POLICY' in line and 'router policy route-map' not in line
                   for line in core.splitlines())
    # ALL also reaches the OOB switch, not merely the core role.
    oob = rendered_sample['oob-switch-01-config.sh']
    assert 'acl example-ssh-allow rule 10' in oob
    assert "system message pre-login 'Example lab - authorized users only'" in oob


def test_route_policy_sample_changes_the_matched_community(rendered_sample):
    core = rendered_sample['core-01-config.sh']
    assert 'router policy community-list EXAMPLE_COMMUNITY rule 10 community 65000:100' in core
    assert 'router policy route-map EXAMPLE_POLICY rule 10 set community 65000:200' in core


def test_banner_sample_explains_and_renders_override_last(rendered_sample):
    wb = openpyxl.load_workbook(ROOT / 'docs/optional-overrides.xlsx', read_only=True)
    try:
        readme = ' '.join(str(row[0]) for row in wb['Read Me'].values)
        assert 'overrides' in readme and 'pre_login_message' in readme
    finally:
        wb.close()
    for name in ('core-01-config.sh', 'oob-switch-01-config.sh'):
        commands = [line for line in rendered_sample[name].splitlines()
                    if line.startswith('nv set system message pre-login ')]
        assert len(commands) >= 2
        assert commands[-1] == "nv set system message pre-login 'Example lab - authorized users only'"
