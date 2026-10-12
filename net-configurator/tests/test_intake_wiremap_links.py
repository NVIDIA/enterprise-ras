# SPDX-FileCopyrightText: Copyright (c) 2026 NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: MIT
"""A Wiremap sheet missing (or renaming) System Name (A)/(Z) must raise,
not silently return an empty (or half-empty) links list.

An empty `links` list makes derive_gpu_planes()/derive_convergence()
return 1/'full' -- turning a real 8-plane dedicated-GPU fabric into a
silently wrong single-plane one. intake_derive's own design rule is that
ambiguity raises rather than defaults; this locks the same rule in at the
column-lookup boundary that feeds it.
"""
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).parent.parent / "scripts"))
sys.path.insert(0, str(Path(__file__).parent))

from intake_fixtures import WIREMAP_HEADERS, build_wiremap_sheet  # noqa: E402
from excel_parser import _intake_wiremap_links  # noqa: E402
from intake_derive import DerivationError  # noqa: E402


def _rename_header(ws, original, new_name):
    idx = WIREMAP_HEADERS.index(original) + 1
    ws.cell(row=2, column=idx, value=new_name)


def test_missing_system_name_a_column_raises():
    # A renamed header behaves identically to an absent one for the
    # normalised-name lookup, so renaming stands in for "missing".
    wb = openpyxl.Workbook()
    ws = build_wiremap_sheet(wb, [])
    _rename_header(ws, "System Name (A)", "Node Name (A)")

    with pytest.raises(DerivationError) as exc_info:
        _intake_wiremap_links(ws)
    assert "system name (a)" in str(exc_info.value)


def test_missing_system_name_z_column_raises():
    wb = openpyxl.Workbook()
    ws = build_wiremap_sheet(wb, [])
    _rename_header(ws, "System Name (Z)", "Node Name (Z)")

    with pytest.raises(DerivationError) as exc_info:
        _intake_wiremap_links(ws)
    assert "system name (z)" in str(exc_info.value)


def test_normal_wiremap_sheet_still_reads_links():
    # Guard against the fix breaking the happy path it must not touch.
    wb = openpyxl.Workbook()
    row = [None] * len(WIREMAP_HEADERS)
    row[3] = "gsl-plane1-01"   # System Name (A)
    row[12] = "core-01"        # System Name (Z)
    ws = build_wiremap_sheet(wb, [row])

    links = _intake_wiremap_links(ws)
    assert links == [{"system_name_a": "gsl-plane1-01", "system_name_z": "core-01"}]


def test_full_chain_one_unrecognised_banding_label_raises_rather_than_deriving_a_wrong_fabric():
    # Reproduces the chain a review traced end to end: one unrecognised
    # group label added to row 1 used to make intake_reader.header_row()
    # fall back to treating row 1 as the real header row, which made
    # column_map() find no 'system name (a)' column, which made
    # _intake_wiremap_links() return an empty links list, which made
    # derive_gpu_planes()/derive_convergence() silently derive an 8-plane
    # dedicated-GPU fabric as a 1-plane converged one -- with no error at
    # any step. Every link in that chain now raises instead; this asserts
    # the chain as a whole does, regardless of which individual link fires.
    wb = openpyxl.Workbook()
    ws = wb.create_sheet("Wiremap")
    ws.cell(1, 1, "Source")
    ws.cell(1, 10, "Destination")
    ws.cell(1, 5, "Extra Band")  # the one label _GROUP_LABELS doesn't know
    for col, name in enumerate(WIREMAP_HEADERS, start=1):
        ws.cell(row=2, column=col, value=name)
    ws.cell(row=3, column=4, value="gsl-plane1-01")   # System Name (A)
    ws.cell(row=3, column=13, value="core-01")        # System Name (Z)

    with pytest.raises(ValueError):  # DerivationError/ValueError, whichever link fires
        _intake_wiremap_links(ws)
