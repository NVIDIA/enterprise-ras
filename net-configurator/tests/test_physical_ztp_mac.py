# SPDX-FileCopyrightText: Copyright (c) 2026 NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: MIT
"""H5 / GitLab #93: a physical switch is identified by its workbook MAC.

`generate_host_vars` gave every switch `generate_mac(name, 'eth0')` and never
read the Switch Data "ZTP MAC Address" column it parses. In Air that is right:
the topology builds each switch VM with that generated eth0 MAC. With
`Deploy in Air = No` the switch is real hardware, so a generated MAC matches
nothing on the wire. The inventory now carries the workbook MAC there, and no
MAC at all for a switch whose row leaves it blank.

Runs the parser against a throwaway copy of the tool tree, like
test_deploy_in_air_local_push.py.
"""
import shutil
import subprocess
import sys
from pathlib import Path

import openpyxl
import yaml

sys.path.insert(0, str(Path(__file__).parent))
import workbook_access  # noqa: E402

NC = Path(__file__).resolve().parent.parent
ARCH = "2-4-3-200"
WORKBOOK_MAC = "48:B0:2D:AA:01:02"
_COPY_IGNORE = shutil.ignore_patterns(
    "output", "archive", ".git", ".venv*", "__pycache__", "*.pyc",
    ".pytest_cache", ".era-secrets", "legacy",
)


def _set_ztp_mac(wb, name, mac):
    ws, cols, first = workbook_access.sheet_table(wb, "Switch Data")
    for r in range(first, ws.max_row + 1):
        if str(ws.cell(r, cols["name"]).value or "").strip() == name:
            ws.cell(r, cols["ztp mac address"]).value = mac
            return
    raise AssertionError(f"no Switch Data row for {name}")


def _generate(tmp_path, deploy_in_air):
    dst = tmp_path / "nc"
    shutil.copytree(NC, dst, ignore=_COPY_IGNORE)
    xlsx = dst / "input" / ARCH / "default" / f"{ARCH}.xlsx"
    wb = openpyxl.load_workbook(xlsx)
    switches = [str(r["name"]) for r in workbook_access.node_rows(wb)
                if r["is_switch"] and str(r["name"]).startswith(("oob-switch", "core-"))]
    if not deploy_in_air:
        workbook_access.set_setting(wb, "deploy_in_air", "No")
        for i, name in enumerate(switches):
            workbook_access.set_node_mgmt_ip(wb, name, f"10.10.0.{11 + i}")
    _set_ztp_mac(wb, switches[0], WORKBOOK_MAC)
    wb.save(xlsx)
    r = subprocess.run([sys.executable, "scripts/excel_parser.py", "--arch", ARCH,
                        "--site", "default", "--skip-validate"],
                       cwd=dst, capture_output=True, text=True)
    assert r.returncode == 0, r.stderr[-2000:]
    hv = dst / "output" / ARCH / "default" / "inventory" / "host_vars"
    return switches, {s: yaml.safe_load((hv / f"{s}.yml").read_text()) for s in switches}


def test_physical_switch_uses_its_workbook_mac(tmp_path):
    switches, hv = _generate(tmp_path, deploy_in_air=False)
    assert hv[switches[0]]["mac_address"] == WORKBOOK_MAC.lower()
    for s in switches[1:]:
        assert "mac_address" not in hv[s], f"{s} has no workbook MAC, so it must not get a made-up one"


def test_air_switch_keeps_the_topology_mac(tmp_path):
    sys.path.insert(0, str(NC / "scripts"))
    from utils import generate_mac
    switches, hv = _generate(tmp_path, deploy_in_air=True)
    for s in switches:
        assert hv[s]["mac_address"] == generate_mac(s, "eth0")
