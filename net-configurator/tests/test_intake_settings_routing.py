# SPDX-FileCopyrightText: Copyright (c) 2026 NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: MIT
"""_parse_intake_settings()'s Deployment routing: the version-role count
assert (a malformed sheet must fail immediately, not confusingly further
downstream) and the derived-overrides-declared log (ADR-0064 says derived
always wins; nothing should say so silently)."""
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).parent.parent / "scripts"))
sys.path.insert(0, str(Path(__file__).parent))

from intake_fixtures import WIREMAP_HEADERS, build_wiremap_sheet  # noqa: E402
from excel_parser import _parse_intake_settings  # noqa: E402


def _deployment_sheet(wb, rows):
    ws = wb.create_sheet("Deployment")
    ws.append(["Setting", "Value"])
    for row in rows:
        ws.append(list(row))
    return ws


def test_raises_when_version_role_count_is_not_four():
    # Routing is shape-based (by name), not position-based: a Deployment
    # sheet missing one of the four Cumulus role versions is
    # indistinguishable, at parse time, from a real setting that collided
    # with a reserved key name and got misfiled as a version. Either way
    # this must fail immediately rather than silently produce a
    # 3-versions dict.
    wb = openpyxl.Workbook()
    build_wiremap_sheet(wb, [])
    _deployment_sheet(wb, [
        ("timezone", "Etc/Zulu"),
        ("core", "5.18.0"),
        ("oob", "5.18.0"),
        ("csl", "5.18.0"),
        # gsl deliberately missing
    ])

    with pytest.raises(ValueError, match="3 Cumulus version row"):
        _parse_intake_settings(wb)


def test_logs_when_a_derived_value_overrides_a_declared_one(capsys):
    wb = openpyxl.Workbook()
    row = [None] * len(WIREMAP_HEADERS)
    row[3] = "gsl-plane1-01"   # System Name (A)
    row[12] = "core-01"        # System Name (Z)
    build_wiremap_sheet(wb, [row])   # derives gpu_planes=1, convergence=dedicated_gpu

    _deployment_sheet(wb, [
        ("gpu_planes", 8),          # declared -- will be overridden
        ("convergence", "full"),    # declared -- will be overridden
        ("core", "5.18.0"), ("oob", "5.18.0"),
        ("csl", "5.18.0"), ("gsl", "5.18.0"),
    ])

    settings, _versions = _parse_intake_settings(wb)

    assert settings["gpu_planes"] == 1
    assert settings["convergence"] == "dedicated_gpu"
    out = capsys.readouterr().out
    assert "gpu_planes" in out and "8" in out
    assert "convergence" in out and "full" in out


def test_no_log_when_nothing_is_declared_to_override(capsys):
    wb = openpyxl.Workbook()
    row = [None] * len(WIREMAP_HEADERS)
    row[3] = "gsl-plane1-01"
    row[12] = "core-01"
    build_wiremap_sheet(wb, [row])

    _deployment_sheet(wb, [
        ("core", "5.18.0"), ("oob", "5.18.0"),
        ("csl", "5.18.0"), ("gsl", "5.18.0"),
    ])

    _parse_intake_settings(wb)

    assert "overrides declared" not in capsys.readouterr().out
