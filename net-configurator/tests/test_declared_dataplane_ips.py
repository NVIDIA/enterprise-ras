# SPDX-FileCopyrightText: Copyright (c) 2026 NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: MIT
"""
GitLab #44 / ERA-43 / ADR-0034 — declared-if-present, allocated-if-blank
data-plane addresses from `Host Data`.

A filled `IP` on a non-`eth0` Host Data row (bond0 / bond1 / GPU NIC) is a
declaration the generator must honour exactly; a blank one is an allocation
request served by the pre-existing allocators. Declarations are validated
against the interface's VLAN, the fabric's gateway/SVI claims, the other rows
of the same bond, and every other address on the fabric -- and every failure
is raised BEFORE a single generated file is written.

Unit tests drive `build_devices(declared_ips=...)` directly with synthetic
nodes; the workbook tests round-trip a shipped default through
`emit_complete_workbook.py`, mutate one declaration and prove the mutated
value -- not a re-allocated one -- lands in the generated inventory.
"""
import importlib.util
import os
import shutil
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest
import yaml

SCRIPTS_DIR = Path(__file__).parent.parent / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

from excel_parser import (  # noqa: E402
    DeclaredIpError,
    _parse_intake_host_declarations,
    build_devices,
)
from intake_layout import host_row_groups  # noqa: E402

NC_ROOT = Path(__file__).parent.parent
ARCH = "2-4-3-200"
DEFAULT_XLSX = NC_ROOT / "input" / ARCH / "default" / f"{ARCH}.xlsx"

VLANS = [
    {'name': 'CPU/In-Band', 'subnet': '172.16.178.0/24', 'gateway': '172.16.178.1'},
    {'name': 'GPU', 'subnet': '192.168.0.0/24', 'gateway': '192.168.0.1'},
    {'name': 'Storage', 'subnet': '172.16.180.0/24', 'gateway': '172.16.180.1'},
    {'name': 'Support', 'subnet': '172.16.179.0/24', 'gateway': '172.16.179.1'},
]


def _node(name, role, i):
    return {'name': name, 'role': role, 'status': 'Active',
            'mgmt_ip': f'192.168.200.{10 + i}', 'mac': f'02:00:00:00:00:{i:02x}'}


def _nodes():
    return [_node('gpu-01', 'GPU', 1), _node('gpu-02', 'GPU', 2),
            _node('stor-01', 'Storage', 3), _node('k8s-01', 'k8s', 4)]


def _decl(logical, ip, row=10, nic='', physical=''):
    return {'row': row, 'logical': logical, 'nic': nic,
            'physical': physical, 'profile': '', 'ip': ip}


def _build(declared):
    return build_devices(_nodes(), VLANS, declared_ips=declared)


def _raises(declared):
    with pytest.raises(DeclaredIpError) as ei:
        _build(declared)
    return str(ei.value)


def _wire(host, role, nic, profile, port):
    """One Wire Map row as excel_parser hands them to build_devices."""
    return {'display_in_air': True, 'system_name': host, 'system_role': role,
            'nic_port': nic, 'net_profile': profile, 'switch_name': 'core-01',
            'switch_role': 'core', 'switch_port': port}


def _wired_nodes():
    return [_node('support-01', 'Support', 1), _node('stor-01', 'Storage', 2),
            _node('k8s-01', 'k8s', 3), _node('bcm-01', 'bcme', 4)]


def _wired_rows():
    """support-01: two Support NICs (bond0 only). stor-01: four Storage NICs
    (bond0 + bond1). k8s-01: four CPU NICs (support falls back to cpu).
    bcm-01: no Wire Map rows at all."""
    rows = [_wire('support-01', 'Support', 'eth1', 'Support', 'swp1'),
            _wire('support-01', 'Support', 'eth2', 'Support', 'swp2')]
    rows += [_wire('stor-01', 'Storage', f'eth{i}', 'Storage', f'swp{10 + i}')
             for i in range(1, 5)]
    rows += [_wire('k8s-01', 'k8s', f'eth{i}', 'CPU/In-Band', f'swp{20 + i}')
             for i in range(1, 5)]
    return rows


def _build_wired(declared):
    return build_devices(_wired_nodes(), VLANS, wiremap_rows=_wired_rows(),
                         declared_ips=declared)


# ---------------------------------------------------------------------------
# Baseline: blank cells allocate exactly as before
# ---------------------------------------------------------------------------

class TestBlankCellsAllocate:
    def test_no_declarations_is_the_legacy_allocation(self):
        d = _build(None)
        assert d['gpu-01']['bond_ip'] == '172.16.178.201/24'
        assert d['gpu-01']['gpu_ips'] == ['192.168.0.50/24', '192.168.0.51/24']
        assert d['gpu-02']['bond_ip'] == '172.16.178.202/24'
        assert d['stor-01']['bond_ip1'] == '172.16.180.101/24'
        assert d['stor-01']['bond_ip2'] == '172.16.180.102/24'
        assert d['k8s-01']['bond_ip1'] == '172.16.179.101/24'

    def test_empty_declaration_map_is_identical_to_none(self):
        assert _build({}) == _build(None)

    def test_a_declaration_on_one_host_does_not_move_the_others(self):
        d = _build({'gpu-01': [_decl('bond0', '172.16.178.240')]})
        assert d['gpu-01']['bond_ip'] == '172.16.178.240/24'
        assert d['gpu-02']['bond_ip'] == '172.16.178.202/24'
        assert d['stor-01']['bond_ip1'] == '172.16.180.101/24'


# ---------------------------------------------------------------------------
# Positive: every slot kind can be declared
# ---------------------------------------------------------------------------

class TestDeclarationsAreHonoured:
    def test_compute_bond0_is_the_cpu_inband_address(self):
        d = _build({'gpu-01': [_decl('bond0', '172.16.178.77')]})
        assert d['gpu-01']['bond_ip'] == '172.16.178.77/24'

    def test_declared_prefix_matching_the_vlan_is_accepted(self):
        d = _build({'gpu-01': [_decl('bond0', '172.16.178.77/24')]})
        assert d['gpu-01']['bond_ip'] == '172.16.178.77/24'

    def test_storage_bond0_and_bond1(self):
        d = _build({'stor-01': [_decl('bond0', '172.16.180.150', row=20),
                                _decl('bond1', '172.16.180.151', row=21)]})
        assert d['stor-01']['bond_ip1'] == '172.16.180.150/24'
        assert d['stor-01']['bond_ip2'] == '172.16.180.151/24'

    def test_storage_one_bond_declared_other_allocated(self):
        d = _build({'stor-01': [_decl('bond1', '172.16.180.151')]})
        assert d['stor-01']['bond_ip1'] == '172.16.180.101/24'
        assert d['stor-01']['bond_ip2'] == '172.16.180.151/24'

    def test_support_pool_bonds(self):
        d = _build({'k8s-01': [_decl('bond0', '172.16.179.60', row=30),
                               _decl('bond1', '172.16.179.61', row=31)]})
        assert d['k8s-01']['bond_ip1'] == '172.16.179.60/24'
        assert d['k8s-01']['bond_ip2'] == '172.16.179.61/24'

    def test_two_bond_member_rows_agreeing_is_fine(self):
        # Host Data is one row per PHYSICAL interface; both members of bond0
        # carry the same IP.
        d = _build({'gpu-01': [_decl('bond0', '172.16.178.77', row=10, physical='eth1'),
                               _decl('bond0', '172.16.178.77', row=11, physical='eth2')]})
        assert d['gpu-01']['bond_ip'] == '172.16.178.77/24'

    def test_gpu_nic_without_a_wire_map_cannot_be_bound(self):
        # GPU NICs bind through the Wire Map's interface order; with no Wire
        # Map there is nothing to bind `eth4` to, so this is refused rather
        # than guessed. The Wire-Map-backed GPU path is covered by the
        # completed-workbook round trip below.
        msg = _raises({'gpu-01': [_decl('eth4', '192.168.0.200', nic='eth4', physical='eth4')]})
        assert 'eth4' in msg and 'cannot bind' in msg


# ---------------------------------------------------------------------------
# Negative: fail closed, all problems at once
# ---------------------------------------------------------------------------

class TestDeclarationsAreValidated:
    def test_malformed_address(self):
        msg = _raises({'gpu-01': [_decl('bond0', '172.16.178.')]})
        assert 'not a valid IPv4 address' in msg

    def test_ipv6_literal_is_not_an_ipv4_address(self):
        msg = _raises({'gpu-01': [_decl('bond0', '2001::1')]})
        assert 'not a valid IPv4 address' in msg

    def test_outside_the_interface_vlan(self):
        msg = _raises({'gpu-01': [_decl('bond0', '10.0.0.5')]})
        assert 'outside the VLAN subnet 172.16.178.0/24' in msg

    def test_wrong_prefix_length(self):
        msg = _raises({'gpu-01': [_decl('bond0', '172.16.178.77/25')]})
        assert 'does not match the VLAN subnet' in msg

    @pytest.mark.parametrize('addr', ['172.16.178.0', '172.16.178.255'])
    def test_network_and_broadcast(self, addr):
        msg = _raises({'gpu-01': [_decl('bond0', addr)]})
        assert 'network/broadcast' in msg

    def test_anycast_gateway(self):
        msg = _raises({'gpu-01': [_decl('bond0', '172.16.178.1')]})
        assert 'VRR/anycast gateway' in msg

    def test_switch_svi(self):
        # ADR-0058: with the default reserve the fabric owns .1 (gateway) plus
        # one SVI per SVI-bearing switch starting at .2.
        msg = _raises({'gpu-01': [_decl('bond0', '172.16.178.2')]})
        assert 'switch SVI' in msg

    def test_bond_members_that_disagree(self):
        msg = _raises({'gpu-01': [_decl('bond0', '172.16.178.77', row=10, physical='eth1'),
                                  _decl('bond0', '172.16.178.78', row=11, physical='eth2')]})
        assert 'row 10' in msg and 'row 11' in msg
        assert 'disagree' in msg or 'conflict' in msg.lower()

    def test_duplicate_across_two_hosts(self):
        msg = _raises({'gpu-01': [_decl('bond0', '172.16.178.77', row=10)],
                       'gpu-02': [_decl('bond0', '172.16.178.77', row=20)]})
        assert 'more than one interface' in msg
        assert 'gpu-01' in msg and 'gpu-02' in msg

    def test_declared_colliding_with_an_allocated_host_is_a_collision(self):
        # gpu-02's blank cell is allocated .202; gpu-01 may not declare it.
        # The allocator is NOT silently re-seated -- the operator is told.
        msg = _raises({'gpu-01': [_decl('bond0', '172.16.178.202', row=10)]})
        assert '172.16.178.202/24' in msg
        assert 'gpu-02' in msg and 'gpu-01' in msg

    def test_unknown_logical_interface(self):
        msg = _raises({'gpu-01': [_decl('bond7', '172.16.178.77')]})
        assert 'bond7' in msg

    def test_declaration_into_an_overflowed_slot_still_faces_uniqueness(self):
        # A /29 CPU subnet allocates gpu-01..03 and leaves gpu-04.. unset
        # (bond_ip None). A declaration on an unset slot must still be
        # checked against every other holder of the address.
        vlans = [dict(v) for v in VLANS]
        vlans[0]['subnet'] = '172.16.178.0/29'
        nodes = [_node(f'gpu-0{i}', 'GPU', i) for i in range(1, 6)]
        base = build_devices(nodes, vlans)
        assert base['gpu-04'].get('bond_ip') is None
        taken = base['gpu-01']['bond_ip']
        with pytest.raises(DeclaredIpError) as ei:
            build_devices(nodes, vlans,
                          declared_ips={'gpu-04': [_decl('bond0', taken.split('/')[0])]})
        assert 'gpu-01' in str(ei.value) and 'gpu-04' in str(ei.value)

    def test_a_second_management_row_with_an_ip_is_reported_not_dropped(self):
        msg = _raises({'gpu-01': [_decl('eth0', '172.16.178.77', row=11)]})
        assert 'row 11' in msg and 'cannot bind' in msg

    def test_unknown_host_is_reported(self):
        msg = _raises({'nowhere-01': [_decl('bond0', '172.16.178.77')]})
        assert 'nowhere-01' in msg

    def test_all_errors_are_collected_in_one_raise(self):
        msg = _raises({'gpu-01': [_decl('bond0', '10.0.0.5', row=10)],
                       'stor-01': [_decl('bond0', '172.16.180.0', row=20)]})
        assert 'row 10' in msg and 'row 20' in msg


# ---------------------------------------------------------------------------
# A declared bond must exist on the host's Wire Map (MR !298 Air review)
# ---------------------------------------------------------------------------

class TestDeclaredBondMustBeWired:
    """The server renderers (roles/*/templates/interfaces.j2, air-deploy.py)
    build bond0 from data NICs [0,1] and bond1 from [2,3]. A declaration on a
    bond the host cannot form used to validate clean and then vanish from
    every generated server config; it must be refused before anything is
    written."""

    def test_allocation_is_unchanged_by_the_wire_map_gate(self):
        d = _build_wired(None)
        assert d['support-01']['interfaces'] == {'support': ['eth1', 'eth2']}
        assert d['stor-01']['interfaces'] == {'storage': ['eth1', 'eth2', 'eth3', 'eth4']}
        assert d['support-01']['bond_ip1'] == '172.16.179.101/24'

    def test_bond1_on_a_two_nic_support_host_is_refused(self):
        # The review's exact repro: support-01 has eth1/eth2 only.
        with pytest.raises(DeclaredIpError) as ei:
            _build_wired({'support-01': [
                _decl('bond1', '172.16.179.240/24', row=112, physical='eth1'),
                _decl('bond1', '172.16.179.240/24', row=113, physical='eth2')]})
        msg = str(ei.value)
        assert 'row 112' in msg and 'row 113' in msg
        assert 'bond1 does not exist' in msg and "['eth1', 'eth2']" in msg

    def test_bond0_on_the_same_two_nic_host_is_honoured(self):
        d = _build_wired({'support-01': [_decl('bond0', '172.16.179.240', row=112)]})
        assert d['support-01']['bond_ip1'] == '172.16.179.240/24'

    def test_dual_bond_storage_host_keeps_both_declarations(self):
        d = _build_wired({'stor-01': [_decl('bond0', '172.16.180.150', row=20),
                                      _decl('bond1', '172.16.180.151', row=21)]})
        assert d['stor-01']['bond_ip1'] == '172.16.180.150/24'
        assert d['stor-01']['bond_ip2'] == '172.16.180.151/24'

    def test_support_pool_host_falls_back_to_cpu_nics_like_the_templates(self):
        d = _build_wired({'k8s-01': [_decl('bond1', '172.16.179.61', row=31)]})
        assert d['k8s-01']['bond_ip2'] == '172.16.179.61/24'

    def test_host_with_no_wire_map_rows_cannot_declare_any_bond(self):
        with pytest.raises(DeclaredIpError) as ei:
            _build_wired({'bcm-01': [_decl('bond0', '172.16.179.70', row=40)]})
        assert 'bond0 does not exist' in ei.value.args[0] and '0 data NIC' in ei.value.args[0]

    def test_unwired_bond_is_reported_alongside_other_faults(self):
        msg = None
        with pytest.raises(DeclaredIpError) as ei:
            _build_wired({'support-01': [_decl('bond1', '172.16.179.240', row=112)],
                          'stor-01': [_decl('bond0', '10.0.0.5', row=20)]})
        msg = str(ei.value)
        assert 'row 112' in msg and 'row 20' in msg

    def test_without_a_wire_map_the_gate_is_not_applied(self):
        # Synthetic callers with no Wire Map (the unit tests above) keep the
        # slot-table behaviour; the gate only has evidence when NICs exist.
        d = _build({'k8s-01': [_decl('bond1', '172.16.179.61')]})
        assert d['k8s-01']['bond_ip2'] == '172.16.179.61/24'


# ---------------------------------------------------------------------------
# Host Data reader
# ---------------------------------------------------------------------------

class TestHostDataReader:
    def _ws(self, rows):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Host Data'
        ws.append(['System Name', 'Physical Interface', 'Logical Interface', 'IP',
                   'NICOSName', 'Mac Address', 'Profile'])
        for r in rows:
            ws.append(r)
        return wb

    def test_eth0_is_management_and_filled_bonds_are_declared(self):
        wb = self._ws([
            ['gpu-01', 'eth0', 'eth0', '192.168.200.10', 'eth0', '', 'OOB'],
            ['gpu-01', 'eth1', 'bond0', '172.16.178.77', 'eth1', '', 'CPU'],
            ['gpu-01', 'eth2', 'bond0', '', 'eth2', '', 'CPU'],
            ['gpu-01', 'eth3', 'eth3', '192.168.0.60', 'eth3', '', 'GPU'],
        ])
        g = host_row_groups(wb['Host Data'])
        assert g['gpu-01']['mgmt'] == 2
        assert [(d['logical'], d['ip'], d['row']) for d in g['gpu-01']['declared']] == [
            ('bond0', '172.16.178.77', 3), ('eth3', '192.168.0.60', 5)]
        decl = _parse_intake_host_declarations(wb)
        assert set(decl) == {'gpu-01'}
        assert decl['gpu-01'][1]['nic'] == 'eth3'

    def test_blank_logical_row_is_the_one_row_per_host_form(self):
        wb = self._ws([['gpu-01', '', '', '192.168.200.10', '', '', '']])
        g = host_row_groups(wb['Host Data'])
        assert g['gpu-01']['mgmt'] == 2 and g['gpu-01']['declared'] == []
        assert _parse_intake_host_declarations(wb) == {}

    def test_second_blank_or_eth0_row_with_ip_becomes_a_declaration(self):
        wb = self._ws([
            ['gpu-01', 'eth0', 'eth0', '192.168.200.10', 'eth0', '', 'OOB'],
            ['gpu-01', 'eth9', '', '172.16.178.77', 'eth9', '', 'CPU'],
            ['gpu-01', 'eth0', 'eth0', '172.16.178.78', 'eth0', '', 'OOB'],
        ])
        g = host_row_groups(wb['Host Data'])
        assert g['gpu-01']['mgmt'] == 2
        assert [(d['logical'], d['ip'], d['row']) for d in g['gpu-01']['declared']] == [
            ('', '172.16.178.77', 3), ('eth0', '172.16.178.78', 4)]

    def test_dataplane_rows_without_a_management_row_are_not_promoted(self):
        wb = self._ws([['gpu-01', 'eth1', 'bond0', '172.16.178.77', 'eth1', '', 'CPU']])
        g = host_row_groups(wb['Host Data'])
        assert g['gpu-01']['mgmt'] is None
        assert len(g['gpu-01']['declared']) == 1

    def test_blank_dataplane_cells_are_not_declarations(self):
        wb = self._ws([
            ['gpu-01', 'eth0', 'eth0', '192.168.200.10', 'eth0', '', 'OOB'],
            ['gpu-01', 'eth1', 'bond0', '', 'eth1', '', 'CPU'],
        ])
        assert _parse_intake_host_declarations(wb) == {}


# ---------------------------------------------------------------------------
# Workbook round trip through the real CLI
# ---------------------------------------------------------------------------

def _run(args, cwd=NC_ROOT):
    env = dict(os.environ)
    env['PATH'] = f"{Path(sys.executable).parent}{os.pathsep}{env.get('PATH', '')}"
    return subprocess.run([sys.executable, *args], cwd=cwd, env=env,
                          capture_output=True, text=True)


@pytest.fixture(scope='module')
def complete_workbook(tmp_path_factory):
    """Emit the completed workbook for the shipped 2-4-3-200 default."""
    if not DEFAULT_XLSX.exists():
        pytest.skip(f"{DEFAULT_XLSX} missing")
    out = tmp_path_factory.mktemp('complete') / f'{ARCH}-complete.xlsx'
    site = f'gl44emit-{os.getpid()}'
    in_dir = NC_ROOT / 'input' / ARCH / site
    in_dir.mkdir(parents=True, exist_ok=True)
    shutil.copy(DEFAULT_XLSX, in_dir / f'{ARCH}.xlsx')
    try:
        r = _run(['scripts/emit_complete_workbook.py', '--source', str(DEFAULT_XLSX),
                  '--arch', ARCH, '--site', site, '--out', str(out)])
    finally:
        shutil.rmtree(in_dir, ignore_errors=True)
        shutil.rmtree(NC_ROOT / 'output' / ARCH / site, ignore_errors=True)
    assert r.returncode == 0, r.stdout + r.stderr
    return out


def _host_rows(ws, hdr, host, logical):
    return [r for r in range(2, ws.max_row + 1)
            if ws.cell(r, hdr['system name']).value == host
            and str(ws.cell(r, hdr['logical interface']).value or '').lower() == logical]


def _load_air_deploy():
    spec = importlib.util.spec_from_file_location(
        'air_deploy_gl44', NC_ROOT / 'scripts' / 'air-deploy.py')
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def _first_rows(path):
    """(bond0 row, gpu row) of the first compute host in a completed Host Data."""
    wb = openpyxl.load_workbook(path)
    ws = wb['Host Data']
    hdr = {str(c.value).strip().lower(): i + 1
           for i, c in enumerate(ws[1]) if c.value}
    bond = gpu = None
    for r in range(2, ws.max_row + 1):
        logical = str(ws.cell(r, hdr['logical interface']).value or '').lower()
        ip = str(ws.cell(r, hdr['ip']).value or '')
        if bond is None and logical == 'bond0' and ip.startswith('172.16.178.'):
            bond = r
        if gpu is None and logical.startswith('eth') and logical != 'eth0' \
                and ip.startswith('192.168.0.'):
            gpu = r
        if bond and gpu:
            break
    return wb, ws, hdr, bond, gpu


class TestCompletedWorkbookRoundTrip:
    def _site(self, tmp_path, name):
        site = f'gl44test-{name}-{os.getpid()}'
        in_dir = NC_ROOT / 'input' / ARCH / site
        out_dir = NC_ROOT / 'output' / ARCH / site
        for d in (in_dir, out_dir):
            shutil.rmtree(d, ignore_errors=True)
        in_dir.mkdir(parents=True)
        yield site, in_dir, out_dir
        for d in (in_dir, out_dir):
            shutil.rmtree(d, ignore_errors=True)

    def test_completed_workbook_validates_clean(self, complete_workbook):
        r = _run(['scripts/validate_excel.py', str(complete_workbook)])
        assert r.returncode == 0, r.stdout + r.stderr

    def test_mutated_declaration_lands_in_generated_inventory(self, complete_workbook, tmp_path):
        gen = self._site(tmp_path, 'mut')
        site, in_dir, out_dir = next(gen)
        try:
            wb, ws, hdr, bond, gpu = _first_rows(complete_workbook)
            assert bond and gpu
            host = ws.cell(bond, hdr['system name']).value
            old_bond = ws.cell(bond, hdr['ip']).value
            old_gpu = ws.cell(gpu, hdr['ip']).value
            new_bond, new_gpu = '172.16.178.241/24', '192.168.0.110/24'
            # every member row of that bond must agree
            for r in range(2, ws.max_row + 1):
                if ws.cell(r, hdr['system name']).value == host and \
                        str(ws.cell(r, hdr['logical interface']).value).lower() == 'bond0':
                    ws.cell(r, hdr['ip']).value = new_bond
            ws.cell(gpu, hdr['ip']).value = new_gpu
            wb.save(in_dir / f'{ARCH}.xlsx')

            r = _run(['scripts/excel_parser.py', '--arch', ARCH, '--site', site])
            assert r.returncode == 0, r.stdout + r.stderr
            assert 'Declared data-plane IPs:' in r.stdout
            text = '\n'.join(p.read_text() for p in (out_dir / 'inventory').rglob('*.yml'))
            assert new_bond in text and new_gpu in text
            assert old_bond not in text, "the allocator's value survived -- declaration ignored"
            assert old_gpu not in text
        finally:
            next(gen, None)

    def test_gpu_declaration_in_the_fabric_range_is_refused(self, complete_workbook, tmp_path):
        gen = self._site(tmp_path, 'gpufab')
        site, in_dir, out_dir = next(gen)
        try:
            wb, ws, hdr, _, gpu = _first_rows(complete_workbook)
            ws.cell(gpu, hdr['ip']).value = '192.168.0.20/24'
            xlsx = in_dir / f'{ARCH}.xlsx'
            wb.save(xlsx)
            r = _run(['scripts/excel_parser.py', '--arch', ARCH, '--site', site, '--skip-validate'])
            assert r.returncode != 0
            assert 'GPU fabric range' in r.stderr
            assert not out_dir.exists()
            v = _run(['scripts/validate_excel.py', str(xlsx)])
            assert v.returncode != 0 and 'GPU fabric range' in v.stdout + v.stderr
        finally:
            next(gen, None)

    def test_unwired_bond1_on_support_is_refused_end_to_end(self, complete_workbook, tmp_path):
        """MR !298 Air review repro: support-01 (eth1/eth2 -> bond0 only) with
        both member rows re-labelled bond1 + a new IP. Validation, the parser
        and the server netplan renderer are all exercised: the renderer
        proves WHY the gate exists (bond1 is never built from two NICs), the
        validator and parser prove the declaration no longer slips through."""
        gen = self._site(tmp_path, 'unwired')
        site, in_dir, out_dir = next(gen)
        try:
            wb = openpyxl.load_workbook(complete_workbook)
            ws = wb['Host Data']
            hdr = {str(c.value).strip().lower(): i + 1
                   for i, c in enumerate(ws[1]) if c.value}
            rows = _host_rows(ws, hdr, 'support-01', 'bond0')
            assert len(rows) == 2, 'fixture drift: support-01 should have exactly two bond0 rows'
            nics = sorted(str(ws.cell(r, hdr['nicosname']).value) for r in rows)
            assert nics == ['eth1', 'eth2']
            declared = '172.16.179.240/24'
            for r in rows:
                ws.cell(r, hdr['logical interface']).value = 'bond1'
                ws.cell(r, hdr['ip']).value = declared
            xlsx = in_dir / f'{ARCH}.xlsx'
            wb.save(xlsx)

            # The artefact-level fact behind the gate: with two data NICs the
            # renderer emits bond0 only, so a bond1 address can never land.
            air = _load_air_deploy()
            netplan = air._render_server_netplan(
                'support-01',
                {'role': 'support', 'eth0_ip': '192.168.200.35',
                 'interfaces': {'support': ['eth0', 'eth1', 'eth2']},
                 'bond_ip1': '172.16.179.101/24', 'bond_ip2': declared},
                {'support_network': '172.16.179.0/24',
                 'support_gateway': '172.16.179.1', 'support_vlan': 400})
            cfg = yaml.safe_load(netplan)
            assert set(cfg['network']['bonds']) == {'bond0'}
            assert declared not in netplan

            v = _run(['scripts/validate_excel.py', str(xlsx)])
            assert v.returncode != 0, 'validate-excel accepted an unbindable bond1 declaration'
            vout = v.stdout + v.stderr
            assert 'bond1 does not exist' in vout and 'support-01' in vout

            r = _run(['scripts/excel_parser.py', '--arch', ARCH, '--site', site, '--skip-validate'])
            assert r.returncode != 0
            assert 'bond1 does not exist' in r.stderr and 'row ' in r.stderr
            assert not out_dir.exists(), 'generation wrote output before failing'
        finally:
            next(gen, None)

    def test_collision_fails_before_anything_is_written(self, complete_workbook, tmp_path):
        gen = self._site(tmp_path, 'col')
        site, in_dir, out_dir = next(gen)
        try:
            wb, ws, hdr, bond, _ = _first_rows(complete_workbook)
            host = ws.cell(bond, hdr['system name']).value
            # steal the NEXT compute host's bond0 address
            victim_ip = None
            for r in range(bond + 1, ws.max_row + 1):
                if ws.cell(r, hdr['system name']).value != host and \
                        str(ws.cell(r, hdr['logical interface']).value).lower() == 'bond0' and \
                        str(ws.cell(r, hdr['ip']).value or '').startswith('172.16.178.'):
                    victim_ip = ws.cell(r, hdr['ip']).value
                    break
            assert victim_ip
            for r in range(2, ws.max_row + 1):
                if ws.cell(r, hdr['system name']).value == host and \
                        str(ws.cell(r, hdr['logical interface']).value).lower() == 'bond0':
                    ws.cell(r, hdr['ip']).value = victim_ip
            xlsx = in_dir / f'{ARCH}.xlsx'
            wb.save(xlsx)

            r = _run(['scripts/excel_parser.py', '--arch', ARCH, '--site', site, '--skip-validate'])
            assert r.returncode != 0
            assert 'DeclaredIpError' in r.stderr or 'cannot be applied' in r.stderr
            assert victim_ip in r.stderr
            assert not out_dir.exists(), "generation wrote output before failing"

            v = _run(['scripts/validate_excel.py', str(xlsx)])
            assert v.returncode != 0
            assert victim_ip in v.stdout + v.stderr
        finally:
            next(gen, None)
