# SPDX-FileCopyrightText: Copyright (c) 2026 NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: MIT
"""Port Side / Cable Split must match the breakout sub-port they describe (GitLab #111).

Sub-port `swpNsM` on a W-way breakout is Port Side `M // (W/2) + 1` and Cable
Split `M % (W/2) + 1`; an unsplit `swpN` has Cable Split `NA`. Both fields are
cabling documentation, so a wrong value is a warning, and generation is not
affected by them (see test_cable_position_does_not_drive_inventory).
"""
import filecmp
import shutil
import sys
from pathlib import Path

import openpyxl
import pytest

NC = Path(__file__).resolve().parent.parent
SCRIPTS_DIR = NC / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

from excel_parser import process_excel_template  # noqa: E402
from validate_excel import ValidationResult, validate_wiremap_cable_position  # noqa: E402

WM_HEADERS = [
    "Display in Air", "System Name (A)", "Port (A)", "Port Side (A)",
    "Cable Split (A)", "System Name (B)", "Port (B)", "Port Side (B)",
    "Cable Split (B)", "Network Profile",
]
PROFILES = {"Two Way": 2, "Four Way": 4, "Eight Way": 8, "Whole": 1}


def _book(rows):
    wb = openpyxl.Workbook()
    vp = wb.active
    vp.title = "VLANs & Profiles"
    vp.append(["Profile"] + [None] * 7 + ["Breakout"])
    for name, breakout in PROFILES.items():
        vp.append([name] + [None] * 7 + [breakout])
    wm = wb.create_sheet("Wire Map")
    wm.append(WM_HEADERS)
    for row in rows:
        wm.append(list(row))
    return wb


def _warnings(rows):
    result = ValidationResult()
    validate_wiremap_cable_position(_book(rows), result)
    return result.warnings


def _host_to(switch_port, side, split, profile):
    return ("Yes", "node-01", "eth0", "NA", "NA", "leaf-01", switch_port, side, split, profile)


@pytest.mark.parametrize("profile,width", [("Two Way", 2), ("Four Way", 4), ("Eight Way", 8)])
def test_correct_positions_pass(profile, width):
    lanes = width // 2
    rows = [_host_to(f"swp1s{m}", m // lanes + 1, str(m % lanes + 1), profile) for m in range(width)]
    assert _warnings(rows) == []


def test_numeric_and_na_cells_are_accepted():
    """openpyxl reads Excel numbers as floats (1.0); a blank or NA end is undocumented, not wrong."""
    assert _warnings([
        _host_to("swp3s1", 1.0, 2.0, "Four Way"),
        _host_to("swp3s2", "NA", "NA", "Four Way"),
        _host_to("swp3s3", None, None, "Four Way"),
        _host_to("swp9", "NA", "NA", "Whole"),
    ]) == []


@pytest.mark.parametrize("port,side,split,profile,expect", [
    ("swp1s1", "1", "1", "Four Way", "Cable Split should be 2"),
    ("swp1s2", "1", "1", "Four Way", "Port Side should be 2"),
    ("swp1s5", "1", "2", "Eight Way", "Port Side should be 2"),
    ("swp1s7", "2", "1", "Eight Way", "Cable Split should be 4"),
    ("swp1s1", "1", "1", "Two Way", "Port Side should be 2"),
    ("swp1s1", "2", "1.5", "Two Way", "Cable Split should be 1"),
    ("swp4", "NA", "1.0", "Whole", "Cable Split should be NA"),
    ("swp4", "2", "NA", "Whole", "Port Side should be NA"),
])
def test_wrong_position_is_reported_with_its_cell(port, side, split, profile, expect):
    warnings = _warnings([_host_to(port, side, split, profile)])
    assert len(warnings) == 1, warnings
    assert expect in warnings[0]
    assert "'Wire Map'!H2" in warnings[0] or "'Wire Map'!I2" in warnings[0]
    assert port in warnings[0]


def test_a_side_switch_end_is_checked_too():
    row = ("Yes", "leaf-01", "swp2s3", "1", "2", "leaf-02", "swp2s3", "2", "2", "Four Way")
    warnings = _warnings([row])
    assert len(warnings) == 1 and "'Wire Map'!D2" in warnings[0], warnings


def test_sub_port_outside_its_breakout_is_left_to_the_breakout_check():
    assert _warnings([_host_to("swp1s5", "1", "1", "Four Way")]) == []


def test_long_lists_are_capped_with_a_count():
    rows = [_host_to(f"swp{p}s1", "1", "1", "Four Way") for p in range(1, 31)]
    warnings = _warnings(rows)
    assert len(warnings) == 21
    assert "10 more Port Side / Cable Split" in warnings[-1]


def _scramble_positions(xlsx):
    wb = openpyxl.load_workbook(xlsx)
    ws = wb["Wiremap"]
    header = [c.value for c in ws[2]]
    cols = [i + 1 for i, h in enumerate(header) if h and h.startswith(("Port Side", "Cable Split"))]
    assert len(cols) == 4, header
    swap = {"1": "2", "2": "1", "3": "4", "4": "3", "NA": "3"}
    changed = 0
    for row in range(3, ws.max_row + 1):
        for col in cols:
            cell = ws.cell(row, col)
            if cell.value in swap:
                cell.value = swap[cell.value]
                changed += 1
    wb.save(xlsx)
    return changed


def test_cable_position_does_not_drive_inventory(tmp_path):
    """Port Side and Cable Split are cabling documentation: no inventory value comes from them."""
    arch = "2-8-9-800"
    shipped = NC / "input" / arch / "default" / f"{arch}.xlsx"
    out = {}
    for name in ("shipped", "scrambled"):
        xlsx = tmp_path / name / f"{arch}.xlsx"
        xlsx.parent.mkdir()
        shutil.copy(shipped, xlsx)
        if name == "scrambled":
            assert _scramble_positions(xlsx) > 1000
        out[name] = tmp_path / name / "inventory"
        process_excel_template(xlsx, out[name], site="default")

    def diff(cmp):
        found = [str(Path(cmp.left) / f) for f in cmp.left_only + cmp.right_only + cmp.diff_files]
        return found + [d for sub in cmp.subdirs.values() for d in diff(sub)]

    files = list(out["shipped"].rglob("*"))
    assert len(files) > 10
    assert diff(filecmp.dircmp(out["shipped"], out["scrambled"])) == []
