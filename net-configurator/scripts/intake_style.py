#!/usr/bin/env python3
# SPDX-FileCopyrightText: Copyright (c) 2026 NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: MIT
"""Visual styling for generated intake-form workbooks (ADR-0064).

A converted workbook is read side by side with the vendor's own template, so
it has to look like it: same blues, same merged group bands, same striped
data rows, comparable column widths.

Everything below was MEASURED from the vendor template with openpyxl
(2026-09-09) and vendored as constants, so conversion has no runtime
dependency on a file that is not in this repo.

Each sheet is laid out DIFFERENTLY in the template -- different header rows,
different fills for the same visual role, merged group bands on some sheets
and not others. An earlier revision applied one uniform style to every sheet;
it produced a flat blue stripe instead of merged group headers and read as
obviously wrong. Hence the per-sheet spec table rather than shared defaults.

Deliberately NOT reproduced: data validation (the locked dropdown lists).
A generated workbook is not hand-typed, and carrying them adds a second place
for the allowed-value list to drift from `validate_excel.py`, which is the
real gate.
"""
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DARK = "FF305496"      # dark blue
MID = "FF4472C4"       # mid blue
TINT = "FFDCE6F1"      # pale blue — data rows, and the Destination band
WHITE = "FFFFFFFF"
TITLE_FONT = "FF1F1F1F"


def _fill(rgb):
    return PatternFill("solid", start_color=rgb, end_color=rgb)


# Per-sheet layout, measured from the template.
#   band:    (row, [(first_col, last_col, label_or_None, fill)]) — merged
#   header:  (row, fill)
# `Host Data` has no band row at all; `Wiremap`'s third group is merged but
# unlabelled in the template, and is reproduced that way.
SHEET_SPECS = {
    "Wiremap": {
        # Spans measured from the template: A1:I1, J1:S1, T1:W1. Note the
        # template puts `Network Profile` (col 19) inside the Destination
        # group, not with the cable columns -- matched deliberately.
        "band": (1, [(1, 9, "Source", MID),
                     (10, 19, "Destination", TINT),
                     (20, 23, None, TINT)]),
        "header": (2, DARK),
    },
    "Switch Data": {
        # Merge spans from the template: A1:C1, D1:E1, G1:M1, N1:R1. The
        # FILL runs wider than the merge for Switch IP's -- F1 is separately
        # coloured in the template so the band reads as one block across
        # D..F. Hence `fill_to`: merge and fill are not the same span.
        "band": (1, [(1, 3, "Switch", DARK),
                     (4, 5, "Switch IP's", DARK, 6),
                     (7, 13, "L3", DARK),
                     (14, 18, "L2", MID)]),
        "header": (2, MID),
    },
    "Network Data": {
        # KNOWN STRUCTURAL DIFFERENCE, not a styling gap: the template opens
        # this sheet with a two-row preamble ("Allocated Subnet", "BGP ASN")
        # and puts the column headers on row 3. We emit headers on row 1.
        # Reproducing the preamble would mean inventing an allocated-subnet
        # value the legacy workbook does not carry, and the converter never
        # invents data. Revisit when something actually reads Network Data.
        "band": None,
        "header": (1, MID),
    },
    "Host Data": {
        "band": None,
        "header": (1, DARK),
    },
}

# Column widths measured from the template, keyed by header text so they
# survive reordering. Anything unlisted is sized from its header.
MEASURED_WIDTHS = {
    "DC Location (A)": 12.1, "System Name (A)": 15.6,
    "Interface Type (A)": 18.0, "Port (A)": 21.1, "Cable Split (A)": 13.8,
    "System Name (Z)": 15.9, "Interface Type (Z)": 15.5, "Port (Z)": 21.0,
    "Network Profile": 14.8,
    "Fabric Name": 25.4, "Use": 27.5, "VLAN": 24.6, "VRF": 16.1,
    "IP Subnet": 32.5, "Gateway Addr": 34.4, "Comments": 22.0,
    "DHCP Relay Client": 16.2,
    "Function": 23.1, "Name": 25.4, "ZTP MAC Address": 26.6,
    "Mgmt IP": 15.9, "Loopback IP": 25.9,
    "System Name": 27.8, "Physical Interface": 24.1,
    "Logical Interface": 15.1, "IP": 20.0, "NICOSName": 19.5,
    "Mac Address": 20.0,
    # Free-text columns: sized for content, not for their short header. A
    # header-length default gave `Notes` 11 chars, which truncates anything
    # anyone would actually write in it. `Comments` keeps its measured 22.0
    # from the template above rather than being widened here.
    "Notes": 44.0, "Description": 36.0,
}

MIN_WIDTH = 11.0
MAX_WIDTH = 46.0


def _width_for(header):
    text = str(header).strip()
    if text in MEASURED_WIDTHS:
        return MEASURED_WIDTHS[text]
    return max(MIN_WIDTH, min(MAX_WIDTH, len(text) + 3.5))


def _apply_band(ws, row, groups):
    """Merged, labelled group bands — the template's most visible feature."""
    for group in groups:
        first, last, label, rgb = group[:4]
        fill_to = group[4] if len(group) > 4 else last
        if last > first:
            ws.merge_cells(start_row=row, start_column=first,
                           end_row=row, end_column=last)
        cell = ws.cell(row, first)
        if label is not None:
            cell.value = label
        cell.fill = _fill(rgb)
        # White reads on the two blues; the pale band needs dark text.
        cell.font = Font(bold=True, size=11,
                         color=WHITE if rgb in (DARK, MID) else TITLE_FONT)
        cell.alignment = Alignment(horizontal="left", vertical="center")
        # Fill the rest of the span so the colour runs the full width. Uses
        # `fill_to`, which may extend past the merge (see Switch Data).
        for c in range(first + 1, fill_to + 1):
            ws.cell(row, c).fill = _fill(rgb)


def _apply_header(ws, row, rgb, n_cols):
    for c in range(1, n_cols + 1):
        cell = ws.cell(row, c)
        cell.fill = _fill(rgb)
        cell.font = Font(bold=True, size=11, color=WHITE)
        # The template does NOT wrap headers — it widens columns instead.
        cell.alignment = Alignment(vertical="center", wrap_text=False)


def _band_rows(ws, first_row, last_row, n_cols):
    """Alternating row stripes, the way the template does it.

    The template tints the FIRST data row and every second row after it,
    leaving the rows between unfilled. Filling every row instead produces a
    flat block that reads as grey and makes long sheets much harder to track
    across -- which is the whole reason the template stripes them.

    Done with explicit per-cell fills rather than an Excel table style: the
    template itself has no table objects (checked), so a styled table would
    also add filter dropdowns and a structural element it does not have.
    """
    pale = _fill(TINT)
    for offset, r in enumerate(range(first_row, last_row + 1)):
        if offset % 2:
            continue                      # leave every second row unfilled
        for c in range(1, n_cols + 1):
            ws.cell(r, c).fill = pale


def _set_widths(ws, headers):
    for idx, header in enumerate(headers, start=1):
        if header is None:
            continue
        ws.column_dimensions[get_column_letter(idx)].width = _width_for(header)


def style_sheet(ws, headers, header_row=None, spec=None, last_row=None):
    """Style a finished sheet to match the template.

    Called after the data is written so the tint covers exactly the rows that
    exist; tinting a fixed range would colour empty rows and make a short
    sheet look padded.
    """
    spec = spec or SHEET_SPECS.get(ws.title)
    n_cols = max(1, len([h for h in headers if h is not None]))

    if spec:
        if spec["band"]:
            band_row, groups = spec["band"]
            _apply_band(ws, band_row, groups)
        header_row, header_fill = spec["header"]
    else:
        header_row = header_row or 1
        header_fill = DARK

    _apply_header(ws, header_row, header_fill, n_cols)
    # Stop where the caller says, not at ws.max_row. On a sheet with stacked
    # sub-tables below the main table, running to the bottom means the main
    # stripe and the section stripe interleave and every row ends up tinted.
    end = ws.max_row if last_row is None else last_row
    if end > header_row:
        _band_rows(ws, header_row + 1, end, n_cols)
    _set_widths(ws, headers)
    ws.freeze_panes = ws.cell(header_row + 1, 1)


def style_read_me(ws, width=112.0):
    """The template's Read Me: one wide column, tinted bold first row."""
    ws.column_dimensions["A"].width = width
    first = ws.cell(1, 1)
    first.fill = _fill(TINT)
    first.font = Font(bold=True, size=11, color=TITLE_FONT)
    for r in range(1, ws.max_row + 1):
        ws.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="top")


def style_stacked_sections(ws, first_scan_row=2):
    """Style the stacked sub-tables that sit below a sheet's main table.

    `Network Data` carries VRFs / Port Profiles / DHCP Relay as separate
    tables further down the same sheet. Styling only the main table left
    those reading as raw unformatted data, and -- because they are WIDER
    than the main table -- their extra columns fell outside the styled
    range entirely.

    A section is a title row (text in column A, nothing in column B)
    followed by a header row (text in both). Detected by shape rather than
    by a hardcoded row number: the sections move as the table above them
    grows, which is the same reason the converter locates them by scanning.
    """
    r = first_scan_row
    while r <= ws.max_row:
        a = ws.cell(r, 1).value
        b = ws.cell(r, 2).value
        if a and not b and ws.cell(r + 1, 1).value and ws.cell(r + 1, 2).value:
            title_row, header_row = r, r + 1
            n_cols = 0
            for c in range(1, ws.max_column + 1):
                if ws.cell(header_row, c).value is not None:
                    n_cols = c
            end = header_row
            while end + 1 <= ws.max_row and ws.cell(end + 1, 1).value:
                end += 1

            title = ws.cell(title_row, 1)
            title.fill = _fill(DARK)
            title.font = Font(bold=True, size=11, color=WHITE)
            for c in range(2, n_cols + 1):
                ws.cell(title_row, c).fill = _fill(DARK)

            _apply_header(ws, header_row, MID, n_cols)
            if end > header_row:
                _band_rows(ws, header_row + 1, end, n_cols)

            _widen_from_headers(ws, header_row, n_cols)
            r = end + 1
            continue
        r += 1


def _widen_from_headers(ws, header_row, n_cols):
    """Widen columns to fit a section's own headers, never narrowing."""
    for c in range(1, n_cols + 1):
        header = ws.cell(header_row, c).value
        if header is None:
            continue
        letter = get_column_letter(c)
        current = ws.column_dimensions[letter].width or 0
        ws.column_dimensions[letter].width = max(current, _width_for(header))


# `Deployment` and `Derived` are ERA additions with no template equivalent,
# so their widths are chosen for their content rather than measured.
KEY_VALUE_WIDTHS = {"Deployment": (34.0, 72.0), "Derived": (24.0, 22.0, 58.0)}

# Settings whose values are prose and must wrap rather than run off the cell
# or be truncated by the next column.
WRAPPED_SETTINGS = {"pre_login_message", "post_login_message"}


def style_key_value_sheet(ws, header_row, banner_names):
    """Style a Setting/Value sheet: section banners, widths, wrapped prose.

    `banner_names` are rows that title a section rather than set a value
    (see intake_reader.SECTION_BANNERS). They get the full-width dark band
    that makes the sheet read as grouped sections instead of one flat list.
    """
    widths = KEY_VALUE_WIDTHS.get(ws.title, (34.0, 72.0))
    for idx, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = w
    n_cols = len(widths)

    _apply_header(ws, header_row, DARK, n_cols)

    stripe = 0
    for r in range(header_row + 1, ws.max_row + 1):
        name = ws.cell(r, 1).value
        text = str(name).strip() if name is not None else ""
        if not text:
            continue
        if text.lower() in banner_names:
            for c in range(1, n_cols + 1):
                ws.cell(r, c).fill = _fill(MID)
            ws.cell(r, 1).font = Font(bold=True, size=11, color=WHITE)
            stripe = 0                       # restart striping inside each section
            continue
        if text.startswith("#"):
            ws.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="top")
            continue
        if stripe % 2 == 0:
            for c in range(1, n_cols + 1):
                ws.cell(r, c).fill = _fill(TINT)
        stripe += 1
        if text in WRAPPED_SETTINGS:
            ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")

    # No freeze panes here. `Deployment` opens with an instruction block, so
    # its header sits ~13 rows down and freezing at it locks a large slab of
    # the sheet in place. These are short key/value sheets that fit on a
    # screen anyway -- freezing buys nothing and gets in the way. The wide
    # data sheets (Wiremap, Host Data, ...) still freeze at their header.


def first_stacked_section_row(ws, first_scan_row=2):
    """Row of the first stacked sub-table title, or None if there is none.

    Same shape rule as `style_stacked_sections`: a title has text in column A
    and nothing in column B, and is followed by a header row with both.
    """
    for r in range(first_scan_row, ws.max_row):
        if (ws.cell(r, 1).value and not ws.cell(r, 2).value
                and ws.cell(r + 1, 1).value and ws.cell(r + 1, 2).value):
            return r
    return None


# Shared by both producers of the intake form: the converter and
# `data-models/generate_arch_excel.py`. Lived in the
# converter until the generator needed it too, and the converter is
# deleted at nc-v10 (ADR-0064 D2) while the styling is not.
import intake_reader  # noqa: E402
from intake_layout import WIREMAP_HEADERS, NETWORK_DATA_HEADERS  # noqa: E402


def style_intake_workbook(out_wb):
    """Match the vendor template's look. See scripts/py.

    Layout differs per sheet, so the per-sheet spec table there is the
    authority; only Deployment/Derived (ERA additions with no template
    equivalent) fall back to a generic header style.
    """
    for ws in out_wb.worksheets:
        if ws.title == "Read Me":
            style_read_me(ws)
        elif ws.title == "Wiremap":
            style_sheet(ws, WIREMAP_HEADERS)
        elif ws.title == "Network Data":
            first_section = first_stacked_section_row(ws)
            main_end = (first_section - 1) if first_section else None
            style_sheet(
                ws, NETWORK_DATA_HEADERS, last_row=main_end)
            # VRFs / Port Profiles / DHCP Relay live below the main table and
            # are WIDER than it -- styling only the main table left them raw
            # and their extra columns outside the styled range entirely.
            style_stacked_sections(ws)
        elif ws.title in ("Switch Data", "Host Data"):
            spec = SHEET_SPECS[ws.title]
            hrow = spec["header"][0]
            headers = [ws.cell(hrow, c).value
                       for c in range(1, ws.max_column + 1)]
            style_sheet(ws, headers)
        elif ws.title in ("Deployment", "Derived"):
            # Deployment opens with `#` orientation notes, so its header is
            # not row 1. Find it rather than assume -- that note block is
            # meant to be edited and will change length.
            header_row = 1
            for r in range(1, min(ws.max_row, 40) + 1):
                first = str(ws.cell(r, 1).value or "").strip()
                if first and not first.startswith("#"):
                    header_row = r
                    break
            style_key_value_sheet(
                ws, header_row, intake_reader.SECTION_BANNERS)
